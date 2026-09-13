"""
FastAPI backend for ScriptGen's web UI - Phase 1 (see README's "Web +
desktop, one codebase" section): login, run a query, edit the grid,
generate the UPDATE/rollback scripts. Talks to SQL Server and builds
scripts through the exact same app.core/app.db modules the Tkinter
desktop app used (script_generator, diff_engine, sql_format, sql_pretty,
mssql) - none of that logic changed, only the UI shell around it did.

Run directly for local dev:
    uvicorn web.server:app --reload --port 8420
Or via the desktop launcher (desktop_launcher.py), which starts this
same app in-process and opens it in a native window instead of a
browser tab - see that file's docstring.
"""
from __future__ import annotations

import copy
import decimal
import os
import re
import sys
import uuid
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, HTTPException, Request, Depends, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, ConfigDict, Field
from starlette.middleware.sessions import SessionMiddleware

from app.config import load_config, save_config, ConnectionConfig
from app.core import ai_assist, diff_engine, schema_check, script_generator, sql_pretty, date_anomaly
from app.core import hierarchy_analysis
from app.core import bulk_checker
from app.core import stats as stats_mod
from app.core import snapshot_diff
from app.db import date_anomaly_history, internal_store, mssql, script_history
from app.db import bulk_checker_db
from web.auth import (
    WebUser, WebUserStore, ROLE_ADMIN, ROLE_EDITOR, ROLE_VIEWER, ROLES, ROLE_RANK,
    consume_first_run_notice, get_session_secret,
)
from web.session_store import sessions, QueryState, date_anomaly_sessions, DateAnomalyState
from web import batch_jobs
from web import menu_access

# Same sys.frozen check app/utils/paths.py uses: a PyInstaller --onefile
# build extracts bundled data (see build_web_desktop.bat's --add-data)
# under sys._MEIPASS at runtime, not next to this source file, so
# __file__-relative lookup alone would 404 every static asset once frozen.
if getattr(sys, "frozen", False):
    STATIC_DIR = Path(sys._MEIPASS) / "web" / "static"  # type: ignore[attr-defined]
else:
    STATIC_DIR = Path(__file__).resolve().parent / "static"

# Process start time + PID, captured once at import time - exposed via the
# unauthenticated /api/server-info route below and shown in the UI (login
# screen + sidebar) so "did my restart actually take effect" is something
# the analyst can SEE instead of guessing. This exists because of a real,
# repeated support cycle: run_web.py binds a FIXED port (8420) and, if that
# port is already open, silently just opens a browser tab to whatever is
# ALREADY running there instead of starting a new server (see run_web.py's
# docstring) - so "closing and re-running run_dev.bat" does nothing at all
# if an old ScriptGen console window/process was still alive in the
# background, and every code change looks like it "didn't take" even though
# it's sitting right there on disk. A visibly stale started_at timestamp is
# the fastest way to catch that instead of re-checking code for the tenth
# time.
SERVER_STARTED_AT = datetime.now(timezone.utc).isoformat()
SERVER_PID = os.getpid()

app = FastAPI(title="ScriptGen Web")
app.add_middleware(
    SessionMiddleware,
    secret_key=get_session_secret(),
    session_cookie="scriptgen_session",
    same_site="lax",
    # https_only left at its default (False) so local/LAN HTTP still
    # works out of the box; turn it on once this sits behind real TLS -
    # see the README's web-deployment note.
)


@app.middleware("http")
async def _no_cache_static_assets(request, call_next):
    """
    This app is under active, frequent development - index.html/app.js/
    styles.css change from one session to the next, and Starlette's
    StaticFiles serves them with default browser-cacheable headers. A
    browser that already has app.js cached will keep running the OLD
    JavaScript after a plain refresh (no full navigation happens in this
    SPA, so a stale cached script silently "does nothing" for whatever
    changed - this bit a real round: new sub-nav tabs rendered from fresh
    HTML but their click handlers, defined in the still-cached old
    app.js, weren't there yet). Blanket no-store on every response is
    fine for a small local/tunnel-only tool like this one - the cost of
    always refetching a few small static files is trivial compared to
    the cost of "why isn't my change showing up" confusion during
    development. Revisit if this ever needs to scale to many concurrent
    users on a slow link.
    """
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    return response


# ---------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------
class LoginRequest(BaseModel):
    username: str
    password: str


def _current_web_user(request: Request) -> WebUser:
    """Loads the full account record for whoever the session cookie says
    is logged in - not just trusting the cookie's username. This means
    deactivating or deleting someone's account takes effect on their
    NEXT request, not just their next login - an admin revoking access
    doesn't have to wait for a session to expire on its own."""
    username = request.session.get("user")
    if not username:
        raise HTTPException(status_code=401, detail="Not logged in.")
    store = WebUserStore.load()
    user = store.get(username)
    if user is None or not user.active:
        request.session.clear()
        raise HTTPException(status_code=401, detail="Your account is no longer active. Please log in again.")
    return user


def require_login(request: Request) -> str:
    return _current_web_user(request).username


def require_editor(request: Request) -> str:
    """Gate for anything that produces/saves a script or otherwise
    changes shared state beyond running a read query - see web/auth.py's
    module docstring for the full role matrix. Returns just the
    username (like require_login) so existing routes swapping to this
    dependency don't need any other code changes."""
    user = _current_web_user(request)
    if ROLE_RANK[user.role] < ROLE_RANK[ROLE_EDITOR]:
        raise HTTPException(status_code=403, detail="This action requires an Editor or Admin account.")
    return user.username


def require_admin(request: Request) -> str:
    """Gate for connections, AI settings, and user management - see
    web/auth.py's module docstring."""
    user = _current_web_user(request)
    if user.role != ROLE_ADMIN:
        raise HTTPException(status_code=403, detail="This action requires an Admin account.")
    return user.username


def get_session_id(request: Request) -> str:
    sid = request.session.get("sid")
    if not sid:
        sid = uuid.uuid4().hex
        request.session["sid"] = sid
    return sid


@app.post("/api/login")
def login(body: LoginRequest, request: Request):
    store = WebUserStore.load()
    if not store.verify(body.username, body.password):
        raise HTTPException(status_code=401, detail="Wrong username or password.")
    user = store.get(body.username)
    request.session["user"] = user.username
    request.session.setdefault("sid", uuid.uuid4().hex)
    consume_first_run_notice()
    return {
        "username": user.username, "role": user.role, "must_change_password": user.must_change_password,
        "allowed_menus": menu_access.MenuAccessStore.load().allowed_for_role(user.role),
    }


@app.post("/api/logout")
def logout(request: Request):
    sid = request.session.get("sid")
    if sid:
        sessions.clear(sid)
        date_anomaly_sessions.clear(sid)
    request.session.clear()
    return {"ok": True}


@app.get("/api/session")
def session_info(request: Request):
    user = _current_web_user(request)
    allowed_menus = menu_access.MenuAccessStore.load().allowed_for_role(user.role)
    return {
        "username": user.username, "role": user.role, "must_change_password": user.must_change_password,
        # Which sidebar pages this role's config currently shows - see
        # web/menu_access.py's module docstring for why this is a
        # visibility convenience, not a second security boundary (the
        # real one stays require_editor/require_admin on each route).
        "allowed_menus": allowed_menus,
    }


@app.get("/api/config/menu-access")
def get_menu_access(user: str = Depends(require_admin)):
    """Admin-only: the full role -> visible-menu-ids map, plus the known
    menu id/label list so the frontend's admin panel can render checkboxes
    without hardcoding the menu list twice (see MENU_LABELS)."""
    store = menu_access.MenuAccessStore.load()
    return {
        "menus": [{"id": m, "label": menu_access.MENU_LABELS[m]} for m in menu_access.MENU_IDS],
        "access": store.access,
        "admin_required_menu": menu_access.ADMIN_REQUIRED_MENU,
    }


class MenuAccessUpdateRequest(BaseModel):
    access: dict[str, list[str]]


@app.put("/api/config/menu-access")
def put_menu_access(body: MenuAccessUpdateRequest, user: str = Depends(require_admin)):
    try:
        store = menu_access.set_access(body.access)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"access": store.access}


@app.get("/api/server-info")
def server_info():
    """
    Deliberately unauthenticated (unlike every other route here) - it
    needs to be checkable from the LOGIN screen itself, before signing in,
    since that's the fastest way to tell whether a "restart" actually
    started a new process. See SERVER_STARTED_AT's comment above for why
    this exists: run_web.py's fixed-port design means a restart can
    silently no-op and just reattach to an old still-running instance.
    """
    return {"started_at": SERVER_STARTED_AT, "pid": SERVER_PID}


# ---------------------------------------------------------------------
# Self-service account actions (any logged-in user, own account only)
# ---------------------------------------------------------------------
class ChangePasswordRequest(BaseModel):
    old_password: str
    new_password: str


@app.post("/api/account/change-password")
def change_own_password(body: ChangePasswordRequest, user: str = Depends(require_login)):
    store = WebUserStore.load()
    try:
        store.change_own_password(user, body.old_password, body.new_password)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True}


# ---------------------------------------------------------------------
# User management (admin only) - see web/auth.py's module docstring for
# the role model and the guard rails (can't deactivate/delete/demote
# yourself or the last active admin) enforced inside WebUserStore itself,
# not just here - so the CLI (web/manage_users.py) can't bypass them either.
# ---------------------------------------------------------------------
@app.get("/api/users")
def list_users(admin: str = Depends(require_admin)):
    store = WebUserStore.load()
    return {"users": [u.to_public_dict() for u in store.list_users()]}


class CreateUserRequest(BaseModel):
    username: str
    password: str
    role: str = ROLE_EDITOR


@app.post("/api/users")
def create_user(body: CreateUserRequest, admin: str = Depends(require_admin)):
    store = WebUserStore.load()
    try:
        user = store.add_user(body.username, body.password, role=body.role, created_by=admin)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return user.to_public_dict()


class SetRoleRequest(BaseModel):
    role: str


@app.put("/api/users/{username}/role")
def set_user_role(username: str, body: SetRoleRequest, admin: str = Depends(require_admin)):
    store = WebUserStore.load()
    try:
        store.set_role(username, body.role)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return store.get(username).to_public_dict()


@app.post("/api/users/{username}/deactivate")
def deactivate_user(username: str, admin: str = Depends(require_admin)):
    store = WebUserStore.load()
    try:
        store.deactivate(username, acting_username=admin)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return store.get(username).to_public_dict()


@app.post("/api/users/{username}/reactivate")
def reactivate_user(username: str, admin: str = Depends(require_admin)):
    store = WebUserStore.load()
    try:
        store.reactivate(username)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return store.get(username).to_public_dict()


class ResetPasswordRequest(BaseModel):
    new_password: str


@app.post("/api/users/{username}/reset-password")
def reset_user_password(username: str, body: ResetPasswordRequest, admin: str = Depends(require_admin)):
    store = WebUserStore.load()
    try:
        store.set_password(username, body.new_password)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True}


@app.delete("/api/users/{username}")
def delete_user(username: str, admin: str = Depends(require_admin)):
    store = WebUserStore.load()
    try:
        store.remove_user(username, acting_username=admin)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True}


# ---------------------------------------------------------------------
# Connection status
# ---------------------------------------------------------------------
@app.get("/api/connection/status")
def connection_status(user: str = Depends(require_login)):
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        return {"connected": False, "message": "No connection configured.", "elapsed_ms": 0}
    ok, message, elapsed_ms = mssql.test_connection(conn)
    return {"connected": ok, "message": message, "elapsed_ms": elapsed_ms}


# ---------------------------------------------------------------------
# Query
# ---------------------------------------------------------------------
class RunQueryRequest(BaseModel):
    sql: str


@app.post("/api/query/run")
def run_query(body: RunQueryRequest, user: str = Depends(require_login), sid: str = Depends(get_session_id)):
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    sql = body.sql.strip()
    if not sql:
        raise HTTPException(status_code=400, detail="Query is empty.")

    try:
        result = mssql.run_query(conn, sql)
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    sessions.set(
        sid,
        QueryState(
            sql=sql,
            columns=result.columns,
            original_rows=result.rows,
            source_schema=result.source_schema,
            source_table=result.source_table,
        ),
    )

    config.add_recent_query(sql)
    save_config(config)

    display_rows = [[diff_engine.cell_display(v) for v in row] for row in result.rows]
    return {
        "columns": result.columns,
        "display_rows": display_rows,
        "row_count": result.row_count,
        "elapsed_ms": result.elapsed_ms,
        "source_schema": result.source_schema,
        "source_table": result.source_table,
    }


class FormatQueryRequest(BaseModel):
    sql: str


@app.post("/api/query/format")
def format_query(body: FormatQueryRequest, user: str = Depends(require_login)):
    return {"sql": sql_pretty.prettify_sql(body.sql)}


@app.get("/api/query/recent")
def recent_queries(user: str = Depends(require_login)):
    config = load_config()
    return {"queries": config.recent_queries}


# ---------------------------------------------------------------------
# Key columns
# ---------------------------------------------------------------------
class KeyLookupRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    schema_: str = Field(default="", alias="schema")
    table: str


@app.post("/api/keys/lookup")
def lookup_keys(body: KeyLookupRequest, user: str = Depends(require_login)):
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    try:
        columns = mssql.get_primary_key_columns(conn, body.schema_ or "dbo", body.table)
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return {"key_columns": columns}


# ---------------------------------------------------------------------
# Grid diff / live SQL preview (mirrors MainWindow._refresh_grid_highlights
# + _refresh_sql_preview, just returning JSON instead of painting a
# tksheet widget)
# ---------------------------------------------------------------------
class GridDiffRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    edited_rows: list[list[str]]
    key_columns: list[str] = []
    schema_: str = Field(default="", alias="schema")
    table: str = ""
    program: str = script_generator.DEFAULT_AUDIT_PROGRAM


def _load_state_or_400(sid: str) -> QueryState:
    state = sessions.get(sid)
    if not state.columns:
        raise HTTPException(status_code=400, detail="Run a query first.")
    return state


@app.post("/api/grid/diff")
def grid_diff(body: GridDiffRequest, user: str = Depends(require_login), sid: str = Depends(get_session_id)):
    state = _load_state_or_400(sid)
    try:
        changes = diff_engine.compute_row_changes(
            state.columns, state.original_rows, body.edited_rows, key_columns=body.key_columns
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    table = body.table or state.source_table or ""
    schema = body.schema_ or state.source_schema
    changed_rows = []
    for change in changes:
        preview_line = ""
        if table:
            preview_line = script_generator.build_preview_line(schema, table, change, program=body.program)
        changed_rows.append(
            {
                "row_index": change.row_index,
                "changed_columns": [c.column for c in change.cell_changes],
                "key_is_full_row": change.key_is_full_row,
                "preview_line": preview_line,
            }
        )
    return {"changes": changed_rows}


# ---------------------------------------------------------------------
# Script generation
# ---------------------------------------------------------------------
class GenerateScriptRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    edited_rows: list[list[str]]
    key_columns: list[str] = []
    schema_: str = Field(default="", alias="schema")
    table: str
    program: str = script_generator.DEFAULT_AUDIT_PROGRAM
    kind: str = "update"  # "update" | "rollback"


@app.post("/api/script/generate")
def generate_script(body: GenerateScriptRequest, user: str = Depends(require_editor), sid: str = Depends(get_session_id)):
    state = _load_state_or_400(sid)
    if not body.table.strip():
        raise HTTPException(status_code=400, detail="Target table is required.")
    try:
        changes = diff_engine.compute_row_changes(
            state.columns, state.original_rows, body.edited_rows, key_columns=body.key_columns
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    schema = body.schema_ or state.source_schema
    builder = script_generator.generate_rollback_script if body.kind == "rollback" else script_generator.generate_update_script
    result = builder(schema, body.table, changes, source_sql=state.sql, program=body.program)

    try:
        script_history.record_script(
            load_config().internal_db_path,
            username=user,
            kind=script_history.KIND_ROLLBACK if body.kind == "rollback" else script_history.KIND_UPDATE,
            schema_name=schema or "",
            table_name=body.table,
            program=body.program,
            statement_count=result.statement_count,
            warning_count=result.warning_count,
            sql_text=result.sql_text,
            source=script_history.SOURCE_WEB,
        )
    except Exception:
        # Same stance as the desktop app: history is a record of the
        # event, never a precondition for seeing the script you just
        # generated - never let a logging failure surface as a 500 here.
        pass

    return {
        "sql_text": result.sql_text,
        "statement_count": result.statement_count,
        "warning_count": result.warning_count,
    }


# ---------------------------------------------------------------------
# Script history (read-only view of app/db/script_history.py; every entry
# generated here or by the Tkinter desktop app on the same machine)
# ---------------------------------------------------------------------
@app.get("/api/history")
def history_list(user: str = Depends(require_login), mine_only: bool = True):
    config = load_config()
    entries = script_history.list_history(
        config.internal_db_path, limit=200, username=user if mine_only else None
    )
    return {
        "entries": [
            {
                "id": e.id,
                "created_at_utc": e.created_at_utc,
                "username": e.username,
                "kind": e.kind,
                "schema_name": e.schema_name,
                "table_name": e.table_name,
                "program": e.program,
                "statement_count": e.statement_count,
                "warning_count": e.warning_count,
                "source": e.source,
            }
            for e in entries
        ]
    }


@app.get("/api/history/{entry_id}")
def history_detail(entry_id: int, user: str = Depends(require_login)):
    config = load_config()
    entry = script_history.get_entry(config.internal_db_path, entry_id)
    if entry is None:
        raise HTTPException(status_code=404, detail="History entry not found.")
    return {
        "id": entry.id,
        "created_at_utc": entry.created_at_utc,
        "username": entry.username,
        "kind": entry.kind,
        "schema_name": entry.schema_name,
        "table_name": entry.table_name,
        "program": entry.program,
        "statement_count": entry.statement_count,
        "warning_count": entry.warning_count,
        "source": entry.source,
        "sql_text": entry.sql_text,
    }


# ---------------------------------------------------------------------
# Config: connections (the web equivalent of the desktop Config dialog's
# connection management + Environment switcher, and of Tools > Environment)
# ---------------------------------------------------------------------
_KEY_RE = re.compile(r"[^a-z0-9_]+")


def _slugify_key(name: str, existing: set[str]) -> str:
    base = _KEY_RE.sub("_", name.strip().lower()).strip("_") or "connection"
    key = base
    i = 2
    while key in existing:
        key = f"{base}_{i}"
        i += 1
    return key


def _connection_out(key: str, conn: ConnectionConfig, active_key: str) -> dict:
    return {
        "key": key,
        "name": conn.name,
        "server": conn.server,
        "port": conn.port,
        "database": conn.database,
        "username": conn.username,
        "timeout_seconds": conn.timeout_seconds,
        "notes": conn.notes,
        "is_active": key == active_key,
    }


class ConnectionIn(BaseModel):
    name: str
    server: str
    port: int = 1433
    database: str = ""
    username: str = ""
    password: str = ""  # required on create; blank on update = keep the existing password
    timeout_seconds: int = 10
    notes: str = ""


@app.get("/api/config/connections")
def list_connections(user: str = Depends(require_login)):
    config = load_config()
    return {
        "connections": [
            _connection_out(key, conn, config.active_connection)
            for key, conn in config.connections.items()
        ]
    }


@app.post("/api/config/connections")
def create_connection(body: ConnectionIn, user: str = Depends(require_admin)):
    config = load_config()
    if not body.name.strip():
        raise HTTPException(status_code=400, detail="Connection name is required.")
    key = _slugify_key(body.name, set(config.connections))
    conn = ConnectionConfig(
        name=body.name, server=body.server, port=body.port, database=body.database,
        username=body.username, timeout_seconds=body.timeout_seconds, notes=body.notes,
    )
    conn.set_password(body.password)
    config.add_connection(key, conn)
    if len(config.connections) == 1:
        config.active_connection = key
    save_config(config)
    return _connection_out(key, conn, config.active_connection)


@app.put("/api/config/connections/{key}")
def update_connection(key: str, body: ConnectionIn, user: str = Depends(require_admin)):
    config = load_config()
    if key not in config.connections:
        raise HTTPException(status_code=404, detail="Connection not found.")
    conn = config.connections[key]
    conn.name = body.name
    conn.server = body.server
    conn.port = body.port
    conn.database = body.database
    conn.username = body.username
    conn.timeout_seconds = body.timeout_seconds
    conn.notes = body.notes
    if body.password:
        conn.set_password(body.password)
    save_config(config)
    return _connection_out(key, conn, config.active_connection)


@app.delete("/api/config/connections/{key}")
def delete_connection(key: str, user: str = Depends(require_admin)):
    config = load_config()
    if key not in config.connections:
        raise HTTPException(status_code=404, detail="Connection not found.")
    try:
        config.remove_connection(key)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    save_config(config)
    return {"ok": True}


@app.post("/api/config/connections/{key}/activate")
def activate_connection(key: str, user: str = Depends(require_admin)):
    config = load_config()
    if key not in config.connections:
        raise HTTPException(status_code=404, detail="Connection not found.")
    config.active_connection = key
    save_config(config)
    return {"active_connection": key}


class ConnectionTestRequest(BaseModel):
    key: Optional[str] = None
    # Ad-hoc "Test before you save" fields, used when `key` is omitted -
    # the web equivalent of the desktop Config dialog's Test Connection
    # button working on a not-yet-saved connection.
    name: str = "Test"
    server: str = ""
    port: int = 1433
    database: str = ""
    username: str = ""
    password: str = ""
    timeout_seconds: int = 10


@app.post("/api/config/connections/test")
def test_connection_endpoint(body: ConnectionTestRequest, user: str = Depends(require_admin)):
    config = load_config()
    if body.key:
        if body.key not in config.connections:
            raise HTTPException(status_code=404, detail="Connection not found.")
        conn = config.connections[body.key]
        if body.password:
            # Test a saved connection with a not-yet-saved password
            # override, without mutating the stored one.
            conn = copy.deepcopy(conn)
            conn.set_password(body.password)
    else:
        conn = ConnectionConfig(
            name=body.name, server=body.server, port=body.port, database=body.database,
            username=body.username, timeout_seconds=body.timeout_seconds,
        )
        conn.set_password(body.password)
    ok, message, elapsed_ms = mssql.test_connection(conn)
    return {"connected": ok, "message": message, "elapsed_ms": elapsed_ms}


# ---------------------------------------------------------------------
# Config: AI (Gemini) settings
# ---------------------------------------------------------------------
class AIConfigIn(BaseModel):
    model: str
    api_key: str = ""  # blank = leave the currently-stored key unchanged
    enabled: bool = True


@app.get("/api/config/ai")
def get_ai_config(user: str = Depends(require_login)):
    config = load_config()
    return {
        "model": config.ai.model,
        "enabled": config.ai.enabled,
        "has_key": bool(config.ai.get_api_key()),
    }


@app.post("/api/config/ai")
def set_ai_config(body: AIConfigIn, user: str = Depends(require_admin)):
    config = load_config()
    config.ai.model = body.model.strip() or config.ai.model
    if body.api_key:
        # Saving a NEW key is itself an affirmative "turn this on" - the
        # enabled checkbox reflects the PREVIOUS saved state when the
        # form loads, so requiring it to also be ticked would silently
        # leave AI Assist off the first time someone pastes in a key.
        config.ai.set_api_key(body.api_key)  # AIConfig.set_api_key also flips enabled True
        config.ai.enabled = True
    else:
        config.ai.enabled = body.enabled and bool(config.ai.get_api_key())
    save_config(config)
    return {"model": config.ai.model, "enabled": config.ai.enabled, "has_key": bool(config.ai.get_api_key())}


class AIKeyTestRequest(BaseModel):
    api_key: str = ""  # blank = test the currently-stored key
    model: str = ""    # blank = test the currently-stored model


@app.post("/api/config/ai/test")
def test_ai_key(body: AIKeyTestRequest, user: str = Depends(require_admin)):
    config = load_config()
    api_key = body.api_key or config.ai.get_api_key()
    model = body.model or config.ai.model
    result = ai_assist.test_api_key(api_key, model)
    return {"success": result.success, "text": result.text}


# ---------------------------------------------------------------------
# Saved queries (Tools tab equivalent - distinct from the rolling,
# unnamed Recent-queries menu on the Workspace page)
# ---------------------------------------------------------------------
class SavedQueryIn(BaseModel):
    name: str
    sql: str


@app.get("/api/queries/saved")
def list_saved_queries(user: str = Depends(require_login)):
    config = load_config()
    return {
        "queries": [
            {"name": q.name, "sql": q.sql, "created_at_utc": q.created_at_utc}
            for q in config.saved_queries
        ]
    }


@app.post("/api/queries/saved")
def save_query(body: SavedQueryIn, user: str = Depends(require_editor)):
    if not body.name.strip() or not body.sql.strip():
        raise HTTPException(status_code=400, detail="Both a name and a query are required.")
    config = load_config()
    config.save_query(body.name, body.sql)
    save_config(config)
    return {"ok": True}


@app.delete("/api/queries/saved/{name}")
def delete_saved_query(name: str, user: str = Depends(require_editor)):
    config = load_config()
    config.delete_query(name)
    save_config(config)
    return {"ok": True}


# ---------------------------------------------------------------------
# Schema validation (Tools tab: "Validate Target Schema")
# ---------------------------------------------------------------------
class SchemaValidateRequest(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    schema_: str = Field(default="dbo", alias="schema")
    table: str
    key_columns: list[str] = []


@app.post("/api/schema/validate")
def validate_schema(
    body: SchemaValidateRequest, user: str = Depends(require_login), sid: str = Depends(get_session_id)
):
    state = _load_state_or_400(sid)
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    if not body.table.strip():
        raise HTTPException(status_code=400, detail="Enter a target table first.")
    try:
        table_columns = mssql.get_table_columns(conn, body.schema_ or "dbo", body.table)
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    result = schema_check.validate_columns(state.columns, body.key_columns, table_columns)
    return {
        "table_exists": result.table_exists,
        "missing_columns": result.missing_columns,
        "missing_key_columns": result.missing_key_columns,
        "extra_table_columns": result.extra_table_columns,
        "ok": result.ok,
    }


# ---------------------------------------------------------------------
# Snapshots (Tools tab: Snapshot Diff Viewer + the internal-DB export the
# Workspace page's Export menu also reaches on the desktop app)
# ---------------------------------------------------------------------
class SnapshotExportRequest(BaseModel):
    name: str


@app.post("/api/snapshot/export")
def export_snapshot_endpoint(
    body: SnapshotExportRequest, user: str = Depends(require_editor), sid: str = Depends(get_session_id)
):
    state = _load_state_or_400(sid)
    if not body.name.strip():
        raise HTTPException(status_code=400, detail="Enter a name for this snapshot.")
    config = load_config()
    snapshot_table = internal_store.export_snapshot(
        config.internal_db_path, body.name.strip(), state.columns, state.original_rows,
        source_sql=state.sql, source_table=state.source_table or "",
    )
    return {"snapshot_table": snapshot_table}


@app.get("/api/snapshot/list")
def list_snapshots_endpoint(user: str = Depends(require_login)):
    config = load_config()
    snaps = internal_store.list_snapshots(config.internal_db_path)
    return {
        "snapshots": [
            {
                "snapshot_table": s.snapshot_table,
                "display_name": s.display_name,
                "created_at_utc": s.created_at_utc,
                "source_table": s.source_table,
                "row_count": s.row_count,
            }
            for s in snaps
        ]
    }


@app.delete("/api/snapshot/{snapshot_table}")
def delete_snapshot_endpoint(snapshot_table: str, user: str = Depends(require_editor)):
    config = load_config()
    internal_store.delete_snapshot(config.internal_db_path, snapshot_table)
    return {"ok": True}


class SnapshotDiffRequest(BaseModel):
    snapshot_a: str
    snapshot_b: str
    key_columns: list[str] = []


@app.post("/api/snapshot/diff")
def diff_snapshots_endpoint(body: SnapshotDiffRequest, user: str = Depends(require_login)):
    config = load_config()
    cols_a, rows_a = internal_store.load_snapshot(config.internal_db_path, body.snapshot_a)
    cols_b, rows_b = internal_store.load_snapshot(config.internal_db_path, body.snapshot_b)
    if cols_a != cols_b:
        raise HTTPException(
            status_code=400,
            detail="These two snapshots don't have the same columns - pick two snapshots of the same table/query.",
        )
    unknown = [c for c in body.key_columns if c not in cols_a]
    if unknown:
        raise HTTPException(status_code=400, detail=f"Unknown key column(s): {', '.join(unknown)}")
    result = snapshot_diff.diff_snapshots(cols_a, rows_a, rows_b, body.key_columns)

    def _row_out(r):
        return {
            "key": r.key,
            "status": r.status,
            "cell_changes": [
                {"column": c.column, "old_value": c.old_value, "new_value": c.new_value}
                for c in r.cell_changes
            ],
        }

    return {
        "added": [_row_out(r) for r in result.added],
        "removed": [_row_out(r) for r in result.removed],
        "changed": [_row_out(r) for r in result.changed],
        "unchanged_count": result.unchanged_count,
        "key_is_full_row": result.key_is_full_row,
    }


# ---------------------------------------------------------------------
# AI Assist (Gemini) - mirrors the desktop AI Assist page + the Script
# page's AI Pre-Flight Review button. Every route reads the key/model
# from the saved AI config (Settings page) rather than taking a key over
# the wire, so a key never has to round-trip through the browser again
# once it's saved.
# ---------------------------------------------------------------------
def _ai_creds_or_400() -> tuple[str, str]:
    config = load_config()
    api_key = config.ai.get_api_key()
    if not config.ai.enabled or not api_key:
        raise HTTPException(
            status_code=400, detail="AI Assist isn't configured - add a Gemini API key in Settings."
        )
    return api_key, config.ai.model


class AISqlRequest(BaseModel):
    sql: str
    intent: str = "improve this query"


@app.post("/api/ai/suggest")
def ai_suggest(body: AISqlRequest, user: str = Depends(require_editor)):
    api_key, model = _ai_creds_or_400()
    result = ai_assist.suggest_snippet(api_key, model, body.sql, body.intent)
    return {"success": result.success, "text": result.text, "extracted_sql": result.extracted_sql}


@app.post("/api/ai/optimize")
def ai_optimize(body: AISqlRequest, user: str = Depends(require_editor)):
    api_key, model = _ai_creds_or_400()
    result = ai_assist.optimize_query(api_key, model, body.sql)
    return {"success": result.success, "text": result.text, "extracted_sql": result.extracted_sql}


@app.post("/api/ai/explain")
def ai_explain(body: AISqlRequest, user: str = Depends(require_login)):
    api_key, model = _ai_creds_or_400()
    result = ai_assist.explain_query(api_key, model, body.sql)
    return {"success": result.success, "text": result.text, "extracted_sql": result.extracted_sql}


class AINlWhereRequest(BaseModel):
    request: str
    table_context: str = ""


@app.post("/api/ai/nl_where")
def ai_nl_where(
    body: AINlWhereRequest, user: str = Depends(require_editor), sid: str = Depends(get_session_id)
):
    api_key, model = _ai_creds_or_400()
    state = sessions.get(sid)
    result = ai_assist.suggest_where_clause(api_key, model, state.columns, body.request, body.table_context)
    return {"success": result.success, "text": result.text, "extracted_sql": result.extracted_sql}


class AIReviewRequest(BaseModel):
    sql_text: str


@app.post("/api/ai/review")
def ai_review(body: AIReviewRequest, user: str = Depends(require_editor)):
    api_key, model = _ai_creds_or_400()
    result = ai_assist.review_script(api_key, model, body.sql_text)
    return {"success": result.success, "text": result.text}


class DateAnomalyExplainRowRequest(BaseModel):
    """
    Mirrors the row shape date_anomaly_detect_all returns (see that
    route) - the frontend sends back exactly the row object it already
    has from its last /detect-all response rather than the id alone, so
    this route doesn't need its own DB round-trip just to explain a
    finding the caller already fetched.
    """
    id_item_to_bill: str = ""
    account: str = ""
    supply: str = ""
    offered_service: str = ""
    contract_status: str = ""
    anomalous_type: str = ""
    anomalous_type_code: str = ""
    anomalous_type_description: str = ""
    anomalous_status: str = ""
    anomalous_status_description: str = ""
    item_status: str = ""
    needs_status_advance: bool = False
    detection_date: str = ""


@app.post("/api/date-anomaly/detect-all/explain")
def date_anomaly_detect_all_explain(body: DateAnomalyExplainRowRequest, user: str = Depends(require_login)):
    """
    AI-assisted explanation of ONE Detect All row - require_login (not
    require_editor) since this only reads/explains, it never mutates
    anything, same read-vs-write split as /api/ai/explain for the SQL
    Query page. Deliberately single-row: the frontend button is only
    enabled when exactly one checkbox is selected (see app.js), so this
    route doesn't need to defend against a multi-row payload itself.
    """
    api_key, model = _ai_creds_or_400()
    row_context = "\n".join(
        [
            f"ID_ITEM_TO_BILL: {body.id_item_to_bill}",
            f"Account: {body.account or '(none)'}",
            f"Supply (NISS): {body.supply or '(none)'}",
            f"Offered service: {body.offered_service or '(none)'}",
            f"Contract status: {body.contract_status or '(none)'}",
            f"Anomaly type: {body.anomalous_type_description or body.anomalous_type or '(unknown)'}"
            + (f" (code {body.anomalous_type_code})" if body.anomalous_type_code else ""),
            f"Anomaly status: {body.anomalous_status_description or body.anomalous_status or '(unknown)'}",
            f"Item-to-bill status: {body.item_status or '(unknown)'}",
            f"Needs STATUS advance to {date_anomaly.ITEM_STATUS_TO}: "
            f"{'Yes' if body.needs_status_advance else 'No'}",
            f"Detected: {body.detection_date or '(unknown)'}",
        ]
    )
    result = ai_assist.explain_anomaly_finding(api_key, model, row_context)
    return {"success": result.success, "text": result.text}


# ---------------------------------------------------------------------
# Dashboard - stats/histogram/correlation/trend data for the current
# session's query result. The heavy lifting (bucketing, Pearson r) is
# the exact same app.core.stats module the desktop Dashboard uses; this
# just returns the numbers as JSON instead of a matplotlib Figure -
# rendering (canvas-drawn bars/histograms/scatter) is the frontend's job.
# ---------------------------------------------------------------------
@app.get("/api/dashboard/stats")
def dashboard_stats(user: str = Depends(require_login), sid: str = Depends(get_session_id)):
    state = _load_state_or_400(sid)
    column_stats = stats_mod.compute_column_stats(state.columns, state.original_rows)
    summary = stats_mod.summary_counts(column_stats)
    return {
        "row_count": len(state.original_rows),
        "column_count": len(state.columns),
        "source_table": state.source_table,
        "summary": summary,
        # Row-level data-quality signal alongside the per-column ones below
        # - exact duplicate rows (full-value match), not just duplicate
        # keys. Cheap to compute (one extra pass, no new query) so it's
        # always included rather than gated behind its own endpoint.
        "duplicate_row_count": stats_mod.duplicate_row_count(state.original_rows),
        # Plain-language observations synthesized from column_stats itself
        # (likely id columns, constant columns, high-null columns, numeric
        # columns with outliers) - see data_quality_flags' own docstring.
        "quality_flags": stats_mod.data_quality_flags(column_stats),
        "columns": [
            {
                "name": s.name,
                "kind": s.kind,
                "row_count": s.row_count,
                "non_null_count": s.non_null_count,
                "null_count": s.null_count,
                "distinct_count": s.distinct_count,
                "min_value": s.min_value,
                "max_value": s.max_value,
                "mean_value": s.mean_value,
                "sum_value": s.sum_value,
                "stdev_value": s.stdev_value,
                "median_value": s.median_value,
                "p25_value": s.p25_value,
                "p75_value": s.p75_value,
                "outlier_count": s.outlier_count,
                "unique_ratio": s.unique_ratio,
                "top_values": s.top_values,
            }
            for s in column_stats.values()
        ],
    }


class HistogramRequest(BaseModel):
    column: str
    bins: int = 12


@app.post("/api/dashboard/histogram")
def dashboard_histogram(
    body: HistogramRequest, user: str = Depends(require_login), sid: str = Depends(get_session_id)
):
    state = _load_state_or_400(sid)
    if body.column not in state.columns:
        raise HTTPException(status_code=400, detail="Unknown column.")
    col_index = state.columns.index(body.column)
    labels, counts = stats_mod.numeric_bucket_histogram(state.original_rows, col_index, bins=body.bins)
    return {"labels": labels, "counts": counts}


def _is_numeric_value(v) -> bool:
    return isinstance(v, stats_mod.NUMERIC_TYPES) and not isinstance(v, bool)


class CorrelationRequest(BaseModel):
    column_a: str
    column_b: str


@app.post("/api/dashboard/correlation")
def dashboard_correlation(
    body: CorrelationRequest, user: str = Depends(require_login), sid: str = Depends(get_session_id)
):
    state = _load_state_or_400(sid)
    if body.column_a not in state.columns or body.column_b not in state.columns:
        raise HTTPException(status_code=400, detail="Unknown column.")
    idx_a = state.columns.index(body.column_a)
    idx_b = state.columns.index(body.column_b)
    r = stats_mod.pearson_correlation(state.original_rows, idx_a, idx_b)

    points = [
        [float(row[idx_a]), float(row[idx_b])]
        for row in state.original_rows
        if _is_numeric_value(row[idx_a]) and _is_numeric_value(row[idx_b])
    ]
    max_points = 500
    if len(points) > max_points:
        step = len(points) / max_points
        points = [points[int(i * step)] for i in range(max_points)]
    return {"r": r, "points": points}


@app.get("/api/dashboard/trend")
def dashboard_trend(table: str = "", user: str = Depends(require_login)):
    config = load_config()
    snaps = internal_store.list_snapshots(config.internal_db_path)
    if table:
        snaps = [s for s in snaps if s.source_table == table]
    snaps = sorted(snaps, key=lambda s: s.created_at_utc)
    return {
        "points": [
            {"created_at_utc": s.created_at_utc, "row_count": s.row_count, "display_name": s.display_name}
            for s in snaps
        ]
    }


# ---------------------------------------------------------------------
# Date Anomaly ("Diff Date System" correction workflow) - web port of the
# desktop app's Date Anomaly tab. State lives server-side per session
# (date_anomaly_sessions, same trade-off as QueryState above) rather than
# round-tripping through the browser, since app.core.date_anomaly needs
# real Python-typed id_reading/item/xml values and full XML text to
# build correct SQL literals - see session_store.py's DateAnomalyState.
# ---------------------------------------------------------------------
def _da_col(row: dict, name: str):
    """Case-insensitive dict lookup - mirrors the identical helper in
    app/ui/main_window.py (desktop version of this same page)."""
    for k, v in row.items():
        if k.lower() == name.lower():
            return v
    return None


class DateAnomalyDetectRequest(BaseModel):
    niss: str
    threshold: int = 0


@app.post("/api/date-anomaly/detect")
def date_anomaly_detect(
    body: DateAnomalyDetectRequest, user: str = Depends(require_login), sid: str = Depends(get_session_id)
):
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    niss = body.niss.strip()
    if not niss:
        raise HTTPException(status_code=400, detail="NISS is required.")

    try:
        detect_result = mssql.run_query(conn, date_anomaly.build_detect_query(niss, body.threshold))
        correct_result = mssql.run_query(conn, date_anomaly.build_correct_date_query(niss, body.threshold))
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    rows = [dict(zip(detect_result.columns, r)) for r in detect_result.rows]
    # Analyst request: flag when this NISS's anomalous readings span more
    # than one billing period - computed straight from the rows already
    # fetched above (r.* already carries ID_BILLING_PERIOD), no extra
    # query needed. The frontend colors the whole table orange for this.
    billing_period_count = len({_da_col(r, "ID_BILLING_PERIOD") for r in rows})

    correct_date = None
    correct_date_reading = None
    if correct_result.rows:
        crow = dict(zip(correct_result.columns, correct_result.rows[0]))
        correct_date = _da_col(crow, "READING_DATE")
        correct_date_reading = _da_col(crow, "ID_READING")

    # Best-effort account/offered-service/contract-status lookup for this
    # NISS - purely to enrich the "Analyze with AI" context below (see
    # date_anomaly.build_niss_account_query's docstring). Never blocks
    # Detect: an analyst without a resolving contracted-service row should
    # still see their anomalous readings, just with account left None.
    account = None
    offered_service = None
    contract_status = None
    try:
        account_result = mssql.run_query(conn, date_anomaly.build_niss_account_query(niss))
        if account_result.rows:
            arow = dict(zip(account_result.columns, account_result.rows[0]))
            account = _da_col(arow, "ACCOUNT")
            offered_service = _da_col(arow, "OFFERED_SERVICE")
            contract_status = _da_col(arow, "CONTRACT_STATUS")
    except mssql.ConnectionError_:
        pass

    explanation = date_anomaly.build_case_explanation(
        niss=niss, threshold=body.threshold, anomaly_count=len(rows),
        correct_date=correct_date, correct_date_source_reading=correct_date_reading,
    )

    # Record this analysis pass to history now - Detect is the earliest
    # point "this NISS was analyzed" is true, regardless of whether the
    # analyst goes on to Resolve/Generate. Non-blocking: same stance as
    # every other Script History write in this app.
    history_id = None
    try:
        history_id = date_anomaly_history.record_analysis(
            load_config().internal_db_path,
            username=user, niss=niss, threshold=body.threshold, anomaly_count=len(rows),
            correct_date=correct_date, correct_date_reading=correct_date_reading,
            explanation=explanation, source=date_anomaly_history.SOURCE_WEB,
        )
    except Exception:
        pass

    date_anomaly_sessions.set(
        sid,
        DateAnomalyState(
            niss=niss, threshold=body.threshold, anomaly_rows=rows,
            correct_date=correct_date, correct_date_reading=correct_date_reading,
            history_id=history_id,
            account=account, offered_service=offered_service, contract_status=contract_status,
        ),
    )

    return {
        "rows": [
            {
                "id_reading": diff_engine.cell_display(_da_col(r, "ID_READING")),
                "id_billing_period": diff_engine.cell_display(_da_col(r, "ID_BILLING_PERIOD")),
                "reading_date": diff_engine.cell_display(_da_col(r, "READING_DATE")),
                "read_status": diff_engine.cell_display(_da_col(r, "READ_STATUS")),
            }
            for r in rows
        ],
        "correct_date": diff_engine.cell_display(correct_date) if correct_date is not None else None,
        "correct_date_reading": diff_engine.cell_display(correct_date_reading)
        if correct_date_reading is not None else None,
        "account": account,
        "offered_service": diff_engine.cell_display(offered_service) if offered_service is not None else None,
        "contract_status": diff_engine.cell_display(contract_status) if contract_status is not None else None,
        "billing_period_count": billing_period_count,
        "explanation": explanation,
    }


@app.post("/api/date-anomaly/resolve")
def date_anomaly_resolve(user: str = Depends(require_login), sid: str = Depends(get_session_id)):
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    da_state = date_anomaly_sessions.get(sid)
    if not da_state.anomaly_rows:
        raise HTTPException(status_code=400, detail="Run Detect first.")

    id_readings = [_da_col(r, "ID_READING") for r in da_state.anomaly_rows]
    try:
        # ID_READING -> list[ID_ITEM_TO_BILL]: GCCOM_READINGS_ITEMSTOBILL can
        # have more than one row per reading, so this must accumulate, not
        # overwrite - a plain dict[reading]=item silently dropped every item
        # past the first one for a given reading (the bug behind a NISS
        # coming back with fewer items-to-bill than expected).
        item_map: dict = {}
        item_sql = date_anomaly.build_item_to_bill_query(id_readings)
        if item_sql:
            item_result = mssql.run_query(conn, item_sql)
            for r in item_result.rows:
                row = dict(zip(item_result.columns, r))
                reading_id = _da_col(row, "ID_READING")
                item_id = _da_col(row, "ID_ITEM_TO_BILL")
                if item_id is None:
                    continue
                item_map.setdefault(reading_id, [])
                if item_id not in item_map[reading_id]:
                    item_map[reading_id].append(item_id)

        item_ids = sorted({v for ids in item_map.values() for v in ids}, key=str)
        xml_rows: dict = {}
        xml_sql = date_anomaly.build_xml_lookup_query(item_ids)
        if xml_sql:
            xml_result = mssql.run_query(conn, xml_sql)
            for r in xml_result.rows:
                row = dict(zip(xml_result.columns, r))
                xml_rows[_da_col(row, "ID_XML")] = _da_col(row, date_anomaly.XML_TO_BILL_COLUMN)

        # Part 5 candidates - item-to-bill ids with an OPEN anomaly
        # (ESTAN00009/ESTAN00001) to cancel. build_anomalous_query already
        # filters to open statuses, so every distinct ID_ITEM_TO_BILL that
        # comes back is one that needs a Part 5 statement.
        anomalous_item_ids: list = []
        anomalous_sql = date_anomaly.build_anomalous_query(item_ids)
        if anomalous_sql:
            anomalous_result = mssql.run_query(conn, anomalous_sql)
            seen_anomalous: set = set()
            for r in anomalous_result.rows:
                row = dict(zip(anomalous_result.columns, r))
                aid = _da_col(row, "ID_ITEM_TO_BILL")
                if aid is not None and aid not in seen_anomalous:
                    seen_anomalous.add(aid)
                    anomalous_item_ids.append(aid)

        # Part 3 candidates - item-to-bill ids still at STTOBILL00, to
        # advance to STTOBILL01. build_item_status_query already filters
        # to that status, so every distinct ID_ITEM_TO_BILL that comes
        # back is one that needs a Part 3 statement.
        item_status_ids: list = []
        item_status_sql = date_anomaly.build_item_status_query(item_ids)
        if item_status_sql:
            item_status_result = mssql.run_query(conn, item_status_sql)
            seen_item_status: set = set()
            for r in item_status_result.rows:
                row = dict(zip(item_status_result.columns, r))
                sid_val = _da_col(row, "ID_ITEM_TO_BILL")
                if sid_val is not None and sid_val not in seen_item_status:
                    seen_item_status.add(sid_val)
                    item_status_ids.append(sid_val)

        reading_cols = mssql.get_table_columns(conn, date_anomaly.READING_SCHEMA, date_anomaly.READING_TABLE)
        item_cols = mssql.get_table_columns(conn, date_anomaly.ADMIN_SCHEMA, date_anomaly.ITEMS_TO_BILL_TABLE)
        reading_has_audit = "update_program" in {c.lower() for c in reading_cols}
        item_has_audit = "update_program" in {c.lower() for c in item_cols}
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    da_state.item_to_bill_map = item_map
    da_state.xml_rows = xml_rows
    da_state.anomalous_item_ids = anomalous_item_ids
    da_state.item_status_ids = item_status_ids
    da_state.reading_has_audit = reading_has_audit
    da_state.item_has_audit = item_has_audit
    date_anomaly_sessions.set(sid, da_state)

    linked = sum(1 for ids in item_map.values() if ids)
    total_items = sum(len(ids) for ids in item_map.values())
    distinct_items = len({v for ids in item_map.values() for v in ids})
    explanation = date_anomaly.build_case_explanation(
        niss=da_state.niss, threshold=da_state.threshold, anomaly_count=len(da_state.anomaly_rows),
        correct_date=da_state.correct_date, correct_date_source_reading=da_state.correct_date_reading,
        item_count=distinct_items, xml_count=len(xml_rows),
        item_status_count=len(item_status_ids), anomalous_count=len(anomalous_item_ids),
    )

    if da_state.history_id is not None:
        try:
            date_anomaly_history.update_analysis(
                load_config().internal_db_path, da_state.history_id,
                item_count=distinct_items, xml_count=len(xml_rows),
                item_status_count=len(item_status_ids), anomalous_count=len(anomalous_item_ids),
                explanation=explanation,
            )
        except Exception:
            pass

    return {
        "linked": linked,
        "total_readings": len(da_state.anomaly_rows),
        "distinct_items": distinct_items,
        "total_items": total_items,
        "xml_count": len(xml_rows),
        "anomalous_count": len(anomalous_item_ids),
        "item_status_count": len(item_status_ids),
        "explanation": explanation,
    }


class DateAnomalyGenerateRequest(BaseModel):
    program: str = script_generator.DEFAULT_AUDIT_PROGRAM
    clean: bool = False
    # Analyst request: when this NISS's anomalous readings span more than
    # one billing period, scope the generated script to just the earliest
    # one - see date_anomaly.lowest_billing_period_rows(). No-op (every
    # reading included as before) when there's only one billing period to
    # begin with.
    lowest_billing_period_only: bool = False


@app.post("/api/date-anomaly/generate")
def date_anomaly_generate(
    body: DateAnomalyGenerateRequest, user: str = Depends(require_editor), sid: str = Depends(get_session_id)
):
    da_state = date_anomaly_sessions.get(sid)
    if not da_state.anomaly_rows or da_state.correct_date is None:
        raise HTTPException(status_code=400, detail="Run Detect (and Resolve Bill Links) first.")
    program = body.program.strip()
    if not program:
        raise HTTPException(status_code=400, detail="Jira / Program # is required.")

    # Analyst request - see date_anomaly.lowest_billing_period_rows' own
    # docstring. Computed from the FULL (unfiltered) anomaly_rows first so
    # billing_period_count below always reflects the true total, then
    # scoped_rows narrows what id_readings/orphan_usage_ids are actually
    # derived from when the option is checked. No-op when there's only one
    # billing period to begin with.
    scoped_rows = (
        date_anomaly.lowest_billing_period_rows(da_state.anomaly_rows)
        if body.lowest_billing_period_only else da_state.anomaly_rows
    )
    id_readings = [_da_col(r, "ID_READING") for r in scoped_rows]
    # Part 6 candidates - non-cycle readings (READING_TYPE TIPTL00011) among
    # the already-fetched anomaly rows (all READ_STATUS 6000STSRED already,
    # by build_detect_query's own WHERE clause - see
    # date_anomaly.READING_TYPE_ORPHAN_USAGE for the analyst's rule). No
    # extra query needed - READING_TYPE/READ_STATUS are already columns on
    # every row from Detect. Derived from scoped_rows, same as id_readings
    # above, so a lowest-billing-period-only script never resets an orphan
    # reading that belongs to an excluded, later billing period.
    orphan_usage_ids = [
        _da_col(r, "ID_READING") for r in scoped_rows
        if _da_col(r, "READING_TYPE") == date_anomaly.READING_TYPE_ORPHAN_USAGE
        and _da_col(r, "READ_STATUS") == date_anomaly.READ_STATUS_ANOMALY
    ]
    billing_period_count = len({_da_col(r, "ID_BILLING_PERIOD") for r in da_state.anomaly_rows})
    result = date_anomaly.build_correction_script(
        niss=da_state.niss,
        threshold=da_state.threshold,
        correct_date=da_state.correct_date,
        correct_date_source_reading=da_state.correct_date_reading,
        anomaly_id_readings=id_readings,
        item_to_bill_map=da_state.item_to_bill_map,
        xml_rows=da_state.xml_rows,
        anomalous_item_ids=da_state.anomalous_item_ids,
        item_status_ids=da_state.item_status_ids,
        orphan_usage_id_readings=orphan_usage_ids,
        billing_period_count=billing_period_count,
        scoped_to_lowest_period=body.lowest_billing_period_only,
        reading_has_audit_cols=da_state.reading_has_audit,
        item_has_audit_cols=da_state.item_has_audit,
        program=program,
        clean=body.clean,
    )

    try:
        script_history.record_script(
            load_config().internal_db_path,
            username=user,
            kind=script_history.KIND_DATE_ANOMALY,
            schema_name="",
            table_name=f"NISS {da_state.niss}",
            program=program,
            statement_count=result.statement_count,
            warning_count=result.warning_count,
            sql_text=result.sql_text,
            source=script_history.SOURCE_WEB,
        )
    except Exception:
        # Same stance as everywhere else - history is a record of the
        # event, never a precondition for seeing the script.
        pass

    if da_state.history_id is not None:
        try:
            date_anomaly_history.update_analysis(
                load_config().internal_db_path, da_state.history_id,
                item_count=result.item_count, xml_count=result.xml_count,
                item_status_count=result.item_status_count, anomalous_count=result.anomalous_count,
                generated=True, explanation=result.explanation,
            )
        except Exception:
            pass

    return {
        "sql_text": result.sql_text,
        "reading_count": result.reading_count,
        "item_count": result.item_count,
        "xml_count": result.xml_count,
        "anomalous_count": result.anomalous_count,
        "item_status_count": result.item_status_count,
        "orphan_reading_count": result.orphan_reading_count,
        "billing_period_count": result.billing_period_count,
        "scoped_to_lowest_period": result.scoped_to_lowest_period,
        "explanation": result.explanation,
        "warnings": result.warnings,
    }


@app.post("/api/date-anomaly/explain-ai")
def date_anomaly_explain_ai(user: str = Depends(require_login), sid: str = Depends(get_session_id)):
    """
    "Analyze with AI" for the Single NISS tab - a Gemini second opinion on
    top of the deterministic date_anomaly.build_case_explanation already
    shown automatically after Detect/Resolve (see those routes above).
    Pulls entirely from server-side session state (DateAnomalyState), same
    as generate - no request body needed since the analyst already ran
    Detect (and optionally Resolve) on this NISS before clicking this.
    """
    api_key, model = _ai_creds_or_400()
    da_state = date_anomaly_sessions.get(sid)
    if not da_state.anomaly_rows and da_state.niss == "":
        raise HTTPException(status_code=400, detail="Run Detect first.")

    distinct_items = len({v for ids in da_state.item_to_bill_map.values() for v in ids})
    billing_period_count = len({_da_col(r, "ID_BILLING_PERIOD") for r in da_state.anomaly_rows})
    orphan_count = sum(
        1 for r in da_state.anomaly_rows
        if _da_col(r, "READING_TYPE") == date_anomaly.READING_TYPE_ORPHAN_USAGE
        and _da_col(r, "READ_STATUS") == date_anomaly.READ_STATUS_ANOMALY
    )
    # Per-reading detail (billing period + reading date), not just a count -
    # the AI had nothing but aggregate numbers before, which made for a
    # generic, not-very-useful "second opinion". Capped so a NISS with an
    # unusually large anomaly count doesn't blow up the prompt.
    _MAX_READING_DETAIL_ROWS = 25
    reading_detail_lines = [
        f"  - ID_READING {_da_col(r, 'ID_READING')}: billing period {_da_col(r, 'ID_BILLING_PERIOD')}, "
        f"reading date {diff_engine.cell_display(_da_col(r, 'READING_DATE'))}, "
        f"reading type {_da_col(r, 'READING_TYPE')}"
        for r in da_state.anomaly_rows[:_MAX_READING_DETAIL_ROWS]
    ]
    if len(da_state.anomaly_rows) > _MAX_READING_DETAIL_ROWS:
        reading_detail_lines.append(f"  - ... and {len(da_state.anomaly_rows) - _MAX_READING_DETAIL_ROWS} more")

    case_context = "\n".join(
        [
            f"NISS: {da_state.niss}",
            f"Account: {da_state.account or '(none found)'}",
            f"Offered service: {da_state.offered_service if da_state.offered_service is not None else '(unknown)'}",
            f"Contract status: {da_state.contract_status if da_state.contract_status is not None else '(unknown)'}",
            f"Billing period floor: {da_state.threshold}",
            f"Anomalous readings found: {len(da_state.anomaly_rows)}",
            f"Distinct billing periods among them: {billing_period_count}"
            + (" (spans more than one billing period - worth flagging)" if billing_period_count > 1 else ""),
            "Anomalous reading detail:",
            *(reading_detail_lines or ["  (none)"]),
            f"Correct date found: "
            f"{diff_engine.cell_display(da_state.correct_date) if da_state.correct_date is not None else '(none found)'}"
            + (f" (from ID_READING {da_state.correct_date_reading})" if da_state.correct_date_reading is not None else ""),
            f"Resolved to distinct item-to-bill rows: {distinct_items}"
            if da_state.item_to_bill_map else "Resolve Bill Links not run yet",
            f"Items needing STATUS advance (STTOBILL00 -> STTOBILL01): {len(da_state.item_status_ids)}",
            f"Open GCCOM_ANOMALOUS records to cancel: {len(da_state.anomalous_item_ids)}",
            f"Non-cycle (READING_TYPE {date_anomaly.READING_TYPE_ORPHAN_USAGE}) readings needing "
            f"GCCOM_READINGS_ITEMSTOBILL delete + READ_STATUS reset to {date_anomaly.READ_STATUS_RESET}: {orphan_count}",
            f"READING audit columns present: {'Yes' if da_state.reading_has_audit else 'No'}",
            f"ITEMS_TO_BILL audit columns present: {'Yes' if da_state.item_has_audit else 'No'}",
        ]
    )
    result = ai_assist.explain_single_niss_case(api_key, model, case_context)
    return {"success": result.success, "text": result.text}


@app.get("/api/date-anomaly/history")
def date_anomaly_history_list(user: str = Depends(require_login), mine_only: bool = True):
    """Every NISS analyzed through Detect (single-NISS or batch), most
    recently updated first - same "your own work, or admin oversight"
    mine_only convention as /api/history. Includes analyses that never
    made it to Generate (Detect found nothing, or the analyst stopped
    to review after Resolve) - see app.db.date_anomaly_history's
    module docstring for why that's deliberate."""
    config = load_config()
    entries = date_anomaly_history.list_analyses(
        config.internal_db_path, limit=100, username=user if mine_only else None
    )
    return {
        "entries": [
            {
                "id": e.id,
                "created_at_utc": e.created_at_utc,
                "updated_at_utc": e.updated_at_utc,
                "username": e.username,
                "niss": e.niss,
                "threshold": e.threshold,
                "source": e.source,
                "anomaly_count": e.anomaly_count,
                "correct_date": e.correct_date,
                "correct_date_reading": e.correct_date_reading,
                "item_count": e.item_count,
                "xml_count": e.xml_count,
                "item_status_count": e.item_status_count,
                "anomalous_count": e.anomalous_count,
                "generated": e.generated,
                "explanation": e.explanation,
            }
            for e in entries
        ]
    }


@app.post("/api/date-anomaly/detect-all")
def date_anomaly_detect_all(user: str = Depends(require_login)):
    """
    System-wide "detect all pending anomalies" scan - unlike Detect above,
    takes no NISS/threshold: runs date_anomaly.build_detect_all_anomalies_
    query() directly against GCCOM_ANOMALOUS (LEFT JOINed to GCCOM_ITEMS_
    TO_BILL for each row's current STATUS) and returns every OPEN Diff
    Date System anomaly found. Stateless - no session state is written
    here, unlike Detect/Resolve/Generate; the frontend holds the returned
    rows and the analyst's checkbox selection itself, and POSTs the
    selected ids straight to /detect-all/generate when ready.

    Two defensive additions on top of the original scan (this route has no
    NISS to scope it down, so it's the one Date Anomaly query that could
    otherwise run away on a large production database):

    1. Schema pre-check - date_anomaly.ANOMALOUS_TYPE_COLUMN (ID_PRINCIPAL_
       ANOMALY) is explicitly flagged as unverified against the real
       GCCOM_ANOMALOUS schema (see its definition). Rather than let a wrong
       guess surface as a raw "Invalid column name" DB error the way the
       ANOMALOUS_STATUS bug did in an earlier round, check the column
       actually exists first and fail with a clear, actionable message.
    2. Row cap - build_detect_all_anomalies_query(limit=...) adds a
       TOP (N) clause; comparing the returned row count to that same limit
       tells the caller whether the result was truncated, surfaced to the
       frontend as `possibly_truncated` so the analyst knows the list isn't
       necessarily complete rather than assuming it is.

    Each row also carries a human-readable status/type description (falls
    back to the raw code if the lookup join found nothing) and billing-
    service enrichment (account/supply/offered-service/contract-status)
    from date_anomaly.build_detect_all_anomalies_query's own LEFT JOINs -
    see that function's docstring for the join chain and its "unverified
    against the live schema" caveat. Also carries all_cycle (bool) -
    whether every reading linked to the item is a "cycle" reading type,
    per date_anomaly.CYCLE_READING_TYPES.
    """
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    try:
        anomalous_cols = mssql.get_table_columns(conn, date_anomaly.ADMIN_SCHEMA, date_anomaly.ANOMALOUS_TABLE)
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    if anomalous_cols and date_anomaly.ANOMALOUS_TYPE_COLUMN.lower() not in {c.lower() for c in anomalous_cols}:
        raise HTTPException(
            status_code=502,
            detail=(
                f"Column '{date_anomaly.ANOMALOUS_TYPE_COLUMN}' was not found on "
                f"{date_anomaly.ADMIN_SCHEMA}.{date_anomaly.ANOMALOUS_TABLE}. This column name is an "
                "unverified assumption (see date_anomaly.ANOMALOUS_TYPE_COLUMN) - check the real "
                "anomaly-type column on GCCOM_ANOMALOUS and update it before Detect All can run."
            ),
        )

    limit = date_anomaly.DETECT_ALL_DEFAULT_LIMIT
    try:
        result = mssql.run_query(conn, date_anomaly.build_detect_all_anomalies_query(limit=limit))
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    rows = [dict(zip(result.columns, r)) for r in result.rows]
    return {
        "rows": [
            {
                "id_item_to_bill": diff_engine.cell_display(_da_col(r, "ID_ITEM_TO_BILL")),
                # Codes are still sent alongside their descriptions (not
                # dropped) so the frontend can fall back to the raw code
                # when a lookup join finds nothing - a code with no
                # matching description row shouldn't read as a blank cell.
                "anomalous_status": diff_engine.cell_display(_da_col(r, date_anomaly.ANOMALOUS_STATUS_COLUMN)),
                "anomalous_status_description": diff_engine.cell_display(_da_col(r, "ANOMALOUS_STATUS_DESC")),
                "anomalous_type": diff_engine.cell_display(_da_col(r, date_anomaly.ANOMALOUS_TYPE_COLUMN)),
                "anomalous_type_code": diff_engine.cell_display(_da_col(r, "ANOMALOUS_TYPE_CODE")),
                "anomalous_type_description": diff_engine.cell_display(_da_col(r, "ANOMALOUS_TYPE_DESC")),
                "item_status": diff_engine.cell_display(_da_col(r, "ITEM_STATUS")),
                "needs_status_advance": _da_col(r, "ITEM_STATUS") == date_anomaly.ITEM_STATUS_FROM,
                "account": diff_engine.cell_display(_da_col(r, "ACCOUNT")),
                "supply": diff_engine.cell_display(_da_col(r, "SUPPLY")),
                "offered_service": diff_engine.cell_display(_da_col(r, "OFFERED_SERVICE")),
                "contract_status": diff_engine.cell_display(_da_col(r, "CONTRACT_STATUS")),
                "detection_date": diff_engine.cell_display(_da_col(r, "DETECTION_DATE")),
                # None (not True/False) when the SQL Server bit itself came
                # back NULL - shouldn't happen given the CASE always
                # resolves to 0/1, but _da_col returning None (unknown
                # column) shouldn't silently read as "all cycle".
                "all_cycle": (
                    bool(_da_col(r, "ALL_CYCLE")) if _da_col(r, "ALL_CYCLE") is not None else None
                ),
                # int, not bool - the frontend derives ">1" for orange-row
                # styling but shows the actual count in a tooltip. None only
                # if the column is somehow absent (same "don't guess"
                # stance as all_cycle above) - a real COUNT(...) always
                # comes back as a real int, 0 included, never NULL.
                "billing_period_count": (
                    int(_da_col(r, "BILLING_PERIOD_COUNT"))
                    if _da_col(r, "BILLING_PERIOD_COUNT") is not None else None
                ),
            }
            for r in rows
        ],
        "possibly_truncated": len(rows) >= limit,
        "limit": limit,
    }


class DetectAllExportRow(BaseModel):
    """One row of whatever the Detect All table currently has VISIBLE
    (i.e. filtered/searched) - the frontend sends back exactly the same
    row shape its own CSV export already builds client-side
    (daCleanupExportCsv's `header` list in app.js), so this endpoint
    never re-derives Detect All's own filter logic server-side; it only
    turns already-selected rows into a real .xlsx workbook, which a
    browser can't build on its own without pulling in an external
    library (deliberately avoided - see README's Detect All charts
    section on why this app doesn't introduce a second/third
    JS dependency for something this small)."""

    id_item_to_bill: str = ""
    account: str = ""
    supply: str = ""
    offered_service: str = ""
    contract_status: str = ""
    anomalous_type: str = ""
    anomalous_status: str = ""
    item_status: str = ""
    needs_status_advance: bool = False
    all_cycle: Optional[bool] = None
    billing_period_count: Optional[int] = None


class DetectAllExportXlsxRequest(BaseModel):
    rows: list[DetectAllExportRow]


@app.post("/api/date-anomaly/detect-all/export-xlsx")
def date_anomaly_detect_all_export_xlsx(body: DetectAllExportXlsxRequest, user: str = Depends(require_login)):
    """
    Excel counterpart to the client-side CSV export next to it
    (#da-cleanup-export-csv-btn) - same columns, same "whatever's
    currently visible" semantics, just built as a real .xlsx workbook
    via openpyxl (already a project dependency - see requirements.txt's
    own comment on it: "xlsx export for the results grid / Dashboard
    stats table") since a browser can't produce that binary format on
    its own. require_login only, same as the CSV export effectively
    is (no server round-trip there at all) - this is a read-only export
    of data the caller's own Detect All scan already returned to them.
    """
    if not body.rows:
        raise HTTPException(status_code=400, detail="No rows to export.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Detect All"
    headers = [
        "ID Item To Bill", "Account", "Supply (NISS)", "Offered Service", "Contract Status",
        "Anomaly Type", "Anomaly Status", "Item Status", "Needs Status Advance", "All Cycle",
        "Billing Periods",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in body.rows:
        ws.append([
            r.id_item_to_bill, r.account, r.supply, r.offered_service, r.contract_status,
            r.anomalous_type, r.anomalous_status, r.item_status,
            "Yes" if r.needs_status_advance else "",
            "Yes" if r.all_cycle is True else ("No" if r.all_cycle is False else ""),
            r.billing_period_count if r.billing_period_count is not None else "",
        ])
    widths = [16, 14, 16, 18, 14, 45, 30, 14, 18, 10, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="detect_all_anomalies.xlsx"'},
    )


class DateAnomalyCleanupGenerateRequest(BaseModel):
    item_ids: list[str]
    program: str = script_generator.DEFAULT_AUDIT_PROGRAM
    clean: bool = False


@app.post("/api/date-anomaly/detect-all/generate")
def date_anomaly_detect_all_generate(
    body: DateAnomalyCleanupGenerateRequest, user: str = Depends(require_editor),
):
    """
    Generates a bulk-cleanup script (date_anomaly.build_cleanup_script -
    STATUS advance + GCCOM_ANOMALOUS cancellation ONLY, see that
    function's docstring for why it stops there) for whatever item-to-
    bill ids the analyst checked on the detect-all list. Re-runs
    build_item_status_query on exactly those ids right before generating
    (the "re-verify at generate time" pattern used elsewhere in this
    module) rather than trusting the detect-all list's snapshot, in case
    something changed between the scan and clicking Generate.
    """
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    # ID_ITEM_TO_BILL is a numeric column everywhere else in this module
    # (format_sql_literal leaves ints unquoted/unprefixed) - these ids
    # round-tripped through the browser as display strings (detect-all is
    # the one Date Anomaly workflow where that happens - see the module
    # docstring's typed-values-never-cross-the-browser note for why every
    # OTHER workflow avoids this), so they're converted back to int here
    # rather than left as str, which would silently format them as
    # N'500'-style text literals instead of a plain 500 - technically
    # often still matches via SQL Server's implicit conversion, but
    # inconsistent with every other query this module builds.
    try:
        item_ids = [int(i.strip()) for i in body.item_ids if str(i).strip()]
    except ValueError:
        raise HTTPException(status_code=400, detail="Item-to-bill ids must be numeric.")
    if not item_ids:
        raise HTTPException(status_code=400, detail="Select at least one item to bill.")
    program = body.program.strip()
    if not program:
        raise HTTPException(status_code=400, detail="Jira / Program # is required.")

    try:
        item_status_ids: list = []
        item_status_sql = date_anomaly.build_item_status_query(item_ids)
        if item_status_sql:
            item_status_result = mssql.run_query(conn, item_status_sql)
            seen: set = set()
            for r in item_status_result.rows:
                row = dict(zip(item_status_result.columns, r))
                sid_val = _da_col(row, "ID_ITEM_TO_BILL")
                if sid_val is not None and sid_val not in seen:
                    seen.add(sid_val)
                    item_status_ids.append(sid_val)
        item_cols = mssql.get_table_columns(conn, date_anomaly.ADMIN_SCHEMA, date_anomaly.ITEMS_TO_BILL_TABLE)
        item_has_audit = "update_program" in {c.lower() for c in item_cols}
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    result = date_anomaly.build_cleanup_script(
        item_ids=item_ids, item_status_ids=item_status_ids,
        item_has_audit_cols=item_has_audit, program=program, clean=body.clean,
    )

    try:
        script_history.record_script(
            config.internal_db_path,
            username=user,
            kind=script_history.KIND_DATE_ANOMALY,
            schema_name="",
            table_name=f"Bulk anomaly cleanup ({len(item_ids)} item(s))",
            program=program,
            statement_count=result.statement_count,
            warning_count=result.warning_count,
            sql_text=result.sql_text,
            source=script_history.SOURCE_WEB,
        )
    except Exception:
        pass

    return {
        "sql_text": result.sql_text,
        "item_status_count": result.item_status_count,
        "anomalous_count": result.anomalous_count,
        "warnings": result.warnings,
    }


class DateAnomalyBatchStartRequest(BaseModel):
    niss_list: list[str]
    threshold: int = 0
    program: str = script_generator.DEFAULT_AUDIT_PROGRAM
    clean: bool = False
    # Analyst request - same option as Single NISS's DateAnomalyGenerateRequest,
    # applied per-NISS across the whole batch. Also what Detect All's
    # "Generate Script for Selected" forwards when it delegates here - see
    # app.js's daCleanupGenerateBtn handler.
    lowest_billing_period_only: bool = False


def _batch_job_or_404(job_id: str, user: WebUser) -> batch_jobs.BatchJob:
    """A job is only visible to whoever started it, or an admin - same
    "your own work, or admin oversight" boundary as /api/history's
    mine_only default. Returns 404 rather than 403 for someone else's
    job id so this doesn't confirm/deny that a given job id exists."""
    job = batch_jobs.jobs.get(job_id)
    if job is None or (job.created_by != user.username and user.role != ROLE_ADMIN):
        raise HTTPException(status_code=404, detail="Batch job not found.")
    return job


@app.post("/api/date-anomaly/batch/start")
def date_anomaly_batch_start(body: DateAnomalyBatchStartRequest, user: str = Depends(require_editor)):
    """
    Kicks off the Detect -> Resolve -> Generate pipeline for every NISS
    in body.niss_list, on a background thread (web/batch_jobs.py) -
    returns almost immediately with a job id rather than blocking the
    request for the whole batch, since a long NISS list against a real
    DB could take minutes. Poll GET /api/date-anomaly/batch/{job_id}
    for progress. A failure on one NISS does not stop the rest of the
    batch - see date_anomaly.BatchNissResult / STATUS_ERROR.
    """
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")

    # Dedup while preserving order, in case the textarea has repeats.
    seen: set[str] = set()
    niss_values: list[str] = []
    for raw in body.niss_list:
        n = raw.strip()
        if n and n not in seen:
            seen.add(n)
            niss_values.append(n)
    if not niss_values:
        raise HTTPException(status_code=400, detail="At least one NISS is required.")
    program = body.program.strip()
    if not program:
        raise HTTPException(status_code=400, detail="Jira / Program # is required.")

    job = batch_jobs.start_batch(
        created_by=user, conn_cfg=conn, niss_list=niss_values, threshold=body.threshold, program=program,
        internal_db_path=config.internal_db_path, clean=body.clean,
        lowest_billing_period_only=body.lowest_billing_period_only,
    )
    return {"job_id": job.id, "status": job.status, "niss_total": job.niss_total}


@app.get("/api/date-anomaly/batch/jobs")
def date_anomaly_batch_jobs(request: Request):
    """Recent batch jobs for the CURRENT user only (not admin-wide - see
    _batch_job_or_404's docstring for why single-job lookups make that
    exception for admins; the list view stays scoped to "your own runs"
    same as /api/history's default), so reopening the Batch card after
    a page refresh can reattach to an in-flight or just-finished job
    instead of losing track of it."""
    user = _current_web_user(request)
    jobs = batch_jobs.jobs.list_for_user(user.username)
    return {"jobs": [j.to_public_dict() for j in jobs]}


@app.get("/api/date-anomaly/batch/{job_id}")
def date_anomaly_batch_status(job_id: str, request: Request):
    user = _current_web_user(request)
    job = _batch_job_or_404(job_id, user)
    return job.to_public_dict()


@app.post("/api/date-anomaly/batch/{job_id}/cancel")
def date_anomaly_batch_cancel(job_id: str, request: Request):
    user = _current_web_user(request)
    if ROLE_RANK[user.role] < ROLE_RANK[ROLE_EDITOR]:
        raise HTTPException(status_code=403, detail="This action requires an Editor or Admin account.")
    job = _batch_job_or_404(job_id, user)
    cancelled = batch_jobs.jobs.request_cancel(job_id)
    if not cancelled:
        raise HTTPException(status_code=400, detail=f"Job is already {job.status} - nothing to cancel.")
    return {"status": "cancel_requested"}


@app.post("/api/date-anomaly/batch/{job_id}/explain-ai")
def date_anomaly_batch_explain_ai(job_id: str, request: Request):
    """
    "Analyze with AI" for the Batch tab - a Gemini summary of a finished
    (or in-progress) job's per-NISS results as a whole, since batch has no
    per-row selection the way Detect All does (see ai_assist.
    explain_batch_run's docstring). require_login only (not require_editor)
    - same read-vs-write split as every other AI explain route here.
    """
    user = _current_web_user(request)
    job = _batch_job_or_404(job_id, user)
    api_key, model = _ai_creds_or_400()
    if not job.results:
        raise HTTPException(status_code=400, detail="This batch job has no results yet.")

    status_counts: dict[str, int] = {}
    lines = [
        f"Job: {job.niss_total} NISS submitted, {job.processed} processed, overall status {job.status}.",
        "",
    ]
    for r in job.results:
        status_counts[r["status"]] = status_counts.get(r["status"], 0) + 1
        line = f"- NISS {r['niss']}: {r['status']}"
        if r["status"] == date_anomaly.STATUS_OK:
            line += (
                f" (readings={r['reading_count']}, items={r['item_count']}, "
                f"xml={r['xml_count']}, status_advance={r['item_status_count']}, "
                f"anomalies_cancelled={r['anomalous_count']}, warnings={r['warning_count']})"
            )
        elif r.get("error"):
            line += f" - {r['error']}"
        lines.append(line)
    lines.insert(2, "Status breakdown: " + ", ".join(f"{k}={v}" for k, v in status_counts.items()))
    batch_context = "\n".join(lines)

    result = ai_assist.explain_batch_run(api_key, model, batch_context)
    return {"success": result.success, "text": result.text}


# ---------------------------------------------------------------------
# Hierarchy Analysis - system-wide scan for PRIMARY meter hierarchies
# (GCGT_RE_MEASUREMENT_POINT.IND_DIST_PPAL = 1) with a not-yet-billed
# reading, plus a per-hierarchy drill-down. See app/core/hierarchy_
# analysis.py's module docstring for the full background and the
# "KNOWN ASSUMPTIONS" this was built against - those assumptions were
# confirmed live against the tunnel DB but NOT yet confirmed by the
# analyst's own review, so this whole section should be treated as a
# draft pending their sign-off. Stateless (unlike Date Anomaly, no
# session state is kept here) - every route re-runs its query fresh.
# ---------------------------------------------------------------------
@app.post("/api/hierarchy-analysis/detect")
def hierarchy_analysis_detect(user: str = Depends(require_login)):
    """
    System-wide scan via hierarchy_analysis.build_pending_primaries_
    query - every primary measuring point (MP_TYPE IN PRIMARY_MP_TYPES,
    MP_STATUS IN MP_STATUS_ALLOWED - the analyst's own correction, see
    the module's "KNOWN ASSUMPTIONS" docstring) with at least one
    not-yet-billed Cycle/Distribution reading for a billing period after
    hierarchy_analysis.MIN_BILLING_PERIOD, narrowed to that primary's
    earliest such pending billing period. Each row also carries
    SECONDARY_COUNT (how many children report to that primary) and
    CALC_MODULE_TYPE (GCCOM_CALCULATION_MODULE's label for the primary's
    own ID_CALCULATION_MODULE code). Schema pre-check on MP_TYPE first,
    same defensive pattern as date_anomaly_detect_all's check on
    ANOMALOUS_TYPE_COLUMN - this whole feature hinges on that one column
    existing and meaning what the analyst says it means, so a wrong
    guess should fail with a clear message rather than a raw "Invalid
    column name" DB error.
    """
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    try:
        mp_cols = mssql.get_table_columns(
            conn, hierarchy_analysis.READING_SCHEMA, hierarchy_analysis.MEASUREMENT_POINT_TABLE,
        )
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    if mp_cols and hierarchy_analysis.MP_TYPE_COLUMN.lower() not in {c.lower() for c in mp_cols}:
        raise HTTPException(
            status_code=502,
            detail=(
                f"Column '{hierarchy_analysis.MP_TYPE_COLUMN}' was not found on "
                f"{hierarchy_analysis.READING_SCHEMA}.{hierarchy_analysis.MEASUREMENT_POINT_TABLE}. "
                "This is the column Hierarchy Analysis uses to identify a 'primary' measuring point - "
                "check the real column name and update hierarchy_analysis.MP_TYPE_COLUMN."
            ),
        )

    limit = hierarchy_analysis.HIERARCHY_DEFAULT_LIMIT
    try:
        result = mssql.run_query(conn, hierarchy_analysis.build_pending_primaries_query(limit=limit))
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    rows = [dict(zip(result.columns, r)) for r in result.rows]
    return {
        "rows": [
            {
                "id_measuring_point": diff_engine.cell_display(_da_col(r, "ID_MEASURING_POINT")),
                "id_main_mp": diff_engine.cell_display(_da_col(r, "ID_MAIN_MP")),
                "secondary_count": diff_engine.cell_display(_da_col(r, "SECONDARY_COUNT")),
                # How many of this primary's secondaries are still not
                # sent to bill (READ_STATUS Available/Anomalous, or no
                # in-scope reading at all) - 0 means the primary's own
                # pending reading is the ONLY thing left to fix in this
                # whole hierarchy. Drives the frontend's green highlight
                # + "primary-only" filter (task, 2026-09-11).
                "secondaries_not_sent_count": diff_engine.cell_display(_da_col(r, "SECONDARIES_NOT_SENT_COUNT")),
                "id_calculation_module": diff_engine.cell_display(_da_col(r, "ID_CALCULATION_MODULE")),
                "calc_module_type": diff_engine.cell_display(_da_col(r, "CALC_MODULE_TYPE")),
                "id_sector_supply": diff_engine.cell_display(_da_col(r, "ID_SECTOR_SUPPLY")),
                "id_meter": diff_engine.cell_display(_da_col(r, "ID_METER")),
                "niss": diff_engine.cell_display(_da_col(r, "NISS")),
                "mp_type": diff_engine.cell_display(_da_col(r, "MP_TYPE")),
                "mp_type_desc": diff_engine.cell_display(_da_col(r, "MP_TYPE_DESC")),
                "mp_status": diff_engine.cell_display(_da_col(r, "MP_STATUS")),
                "mp_status_desc": diff_engine.cell_display(_da_col(r, "MP_STATUS_DESC")),
                "perc_dist": diff_engine.cell_display(_da_col(r, "PERC_DIST")),
                "id_billing_period": diff_engine.cell_display(_da_col(r, "ID_BILLING_PERIOD")),
                "billing_period_desc": diff_engine.cell_display(_da_col(r, "BILLING_PERIOD_DESC")),
                "id_reading": diff_engine.cell_display(_da_col(r, "ID_READING")),
                "reading_date": diff_engine.cell_display(_da_col(r, "READING_DATE")),
                "reading_prev_date": diff_engine.cell_display(_da_col(r, "READING_PREV_DATE")),
                "reading_type": diff_engine.cell_display(_da_col(r, "READING_TYPE")),
                "reading_type_desc": diff_engine.cell_display(_da_col(r, "READING_TYPE_DESC")),
                "read_status": diff_engine.cell_display(_da_col(r, "READ_STATUS")),
                "value": diff_engine.cell_display(_da_col(r, "VALUE")),
                # READY_USAGE is the analyst's own most-important field on this
                # page - the finalized/validated usage, distinct from the raw
                # READING_USAGE next to it. Surfaced explicitly rather than
                # left for the generic column dump other fields get.
                "ready_usage": diff_engine.cell_display(_da_col(r, "READY_USAGE")),
                "reading_usage": diff_engine.cell_display(_da_col(r, "READING_USAGE")),
                "usage_type": diff_engine.cell_display(_da_col(r, "USAGE_TYPE")),
                "ind_estimate": diff_engine.cell_display(_da_col(r, "IND_ESTIMATE")),
            }
            for r in rows
        ],
        "possibly_truncated": len(rows) >= limit,
        "limit": limit,
    }


class HierarchyDetailRequest(BaseModel):
    id_measuring_point: str
    # Optional: the pending-primary row's own ID_BILLING_PERIOD. When
    # present, scopes every child's reading to that SAME billing period
    # (the analyst's own request) instead of the child's full reading
    # history across every period. The frontend always sends this now
    # (it has the value on hand from the row that was clicked); left
    # optional here so a direct API call without it still works.
    id_billing_period: str | None = None


@app.post("/api/hierarchy-analysis/detail")
def hierarchy_analysis_detail(body: HierarchyDetailRequest, user: str = Depends(require_login)):
    """
    Drill-down for one hierarchy - hierarchy_analysis.build_hierarchy_
    detail_query, the analyst's own reference query parameterized on the
    id. Opened from a pending-primaries row (id_measuring_point there),
    shows every member of that primary's hierarchy, scoped to that same
    row's billing period when id_billing_period is supplied.
    """
    mp_id = (body.id_measuring_point or "").strip()
    if not mp_id:
        raise HTTPException(status_code=400, detail="id_measuring_point is required.")
    try:
        mp_id_int = int(mp_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="id_measuring_point must be numeric.")

    bp_raw = (body.id_billing_period or "").strip()
    bp_int: int | None = None
    if bp_raw:
        try:
            bp_int = int(bp_raw)
        except ValueError:
            raise HTTPException(status_code=400, detail="id_billing_period must be numeric.")

    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    try:
        result = mssql.run_query(
            conn, hierarchy_analysis.build_hierarchy_detail_query(mp_id_int, bp_int)
        )
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    rows = [dict(zip(result.columns, r)) for r in result.rows]
    return {
        "rows": [
            {
                "id_measuring_point": diff_engine.cell_display(_da_col(r, "ID_MEASURING_POINT")),
                "id_main_mp": diff_engine.cell_display(_da_col(r, "ID_MAIN_MP")),
                # Needed so the frontend's reading-history popup (opened
                # from this row) can call /api/hierarchy-analysis/reading-
                # history without a second round-trip to look it up.
                "id_sector_supply": diff_engine.cell_display(_da_col(r, "ID_SECTOR_SUPPLY")),
                "niss": diff_engine.cell_display(_da_col(r, "NISS")),
                "ind_dist_ppal": diff_engine.cell_display(_da_col(r, "IND_DIST_PPAL")),
                "perc_dist": diff_engine.cell_display(_da_col(r, "PERC_DIST")),
                "mp_type": diff_engine.cell_display(_da_col(r, "MP_TYPE")),
                "status": diff_engine.cell_display(_da_col(r, "STATUS")),
                "id_billing_period": diff_engine.cell_display(_da_col(r, "ID_BILLING_PERIOD")),
                "billing_period_desc": diff_engine.cell_display(_da_col(r, "BILLING_PERIOD_DESC")),
                "id_reading": diff_engine.cell_display(_da_col(r, "ID_READING")),
                "reading_date": diff_engine.cell_display(_da_col(r, "READING_DATE")),
                "reading_prev_date": diff_engine.cell_display(_da_col(r, "READING_PREV_DATE")),
                "reading_type": diff_engine.cell_display(_da_col(r, "READING_TYPE")),
                "reading_type_desc": diff_engine.cell_display(_da_col(r, "READING_TYPE_DESC")),
                "read_status": diff_engine.cell_display(_da_col(r, "READ_STATUS")),
                "value": diff_engine.cell_display(_da_col(r, "VALUE")),
                "ready_usage": diff_engine.cell_display(_da_col(r, "READY_USAGE")),
            }
            for r in rows
        ],
    }


class ReadingHistoryRequest(BaseModel):
    id_sector_supply: str


@app.post("/api/hierarchy-analysis/reading-history")
def hierarchy_analysis_reading_history(body: ReadingHistoryRequest, user: str = Depends(require_login)):
    """
    Reading-history popup (task #143) - hierarchy_analysis.build_reading_
    history_query, the analyst's own draft query adapted per their two
    rounds of review: filters directly on ID_SECTOR_SUPPLY (no join to
    GCCOM_SECTOR_SUPPLY needed), and selects the specific columns they
    asked for. Opened from a Hierarchy Detail row (id_sector_supply is
    already on that row's own response - see hierarchy_analysis_detail
    above). Highlighting "the reading period we're checking" is handled
    entirely client-side in app.js, not here - this route just returns
    every reading for the supply, newest billing period first.
    """
    supply_raw = (body.id_sector_supply or "").strip()
    if not supply_raw:
        raise HTTPException(status_code=400, detail="id_sector_supply is required.")
    try:
        supply_int = int(supply_raw)
    except ValueError:
        raise HTTPException(status_code=400, detail="id_sector_supply must be numeric.")

    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    try:
        result = mssql.run_query(conn, hierarchy_analysis.build_reading_history_query(supply_int))
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    rows = [dict(zip(result.columns, r)) for r in result.rows]
    return {
        "rows": [
            {
                "billing_period": diff_engine.cell_display(_da_col(r, "BILLING_PERIOD")),
                "id_reading": diff_engine.cell_display(_da_col(r, "ID_READING")),
                "reading_type": diff_engine.cell_display(_da_col(r, "READING_TYPE")),
                "usage_type": diff_engine.cell_display(_da_col(r, "USAGE_TYPE")),
                "read_status": diff_engine.cell_display(_da_col(r, "READ_STATUS")),
                "prev_date": diff_engine.cell_display(_da_col(r, "PREV DATE")),
                "reading_date": diff_engine.cell_display(_da_col(r, "READING_DATE")),
                "prev_value": diff_engine.cell_display(_da_col(r, "PREV_VALUE")),
                "value": diff_engine.cell_display(_da_col(r, "VALUE")),
                "reading_usage": diff_engine.cell_display(_da_col(r, "READING_USAGE")),
                "ready_usage": diff_engine.cell_display(_da_col(r, "READY_USAGE")),
                "ind_estimate": diff_engine.cell_display(_da_col(r, "IND_ESTIMATE")),
            }
            for r in rows
        ],
    }


class HierarchyExportRow(BaseModel):
    """Mirrors DetectAllExportRow's rationale - see that class's own
    docstring. Same shape as one row of hierarchy_analysis_detect's own
    response, sent back by the frontend for whatever's currently
    visible/filtered."""

    id_measuring_point: str = ""
    id_main_mp: str = ""
    niss: str = ""
    mp_type: str = ""
    mp_status: str = ""
    secondary_count: str = ""
    # See hierarchy_analysis_detect's own comment on this field - 0 means
    # the primary's own reading is the only thing left to fix in the
    # whole hierarchy (task, 2026-09-11).
    secondaries_not_sent_count: str = ""
    id_calculation_module: str = ""
    calc_module_type: str = ""
    id_billing_period: str = ""
    billing_period_desc: str = ""
    id_reading: str = ""
    reading_date: str = ""
    reading_type: str = ""
    reading_type_desc: str = ""
    read_status: str = ""
    # The analyst's own most-important field on this page - see the
    # matching comment on hierarchy_analysis_detect's response mapping.
    ready_usage: str = ""
    reading_usage: str = ""


class HierarchyExportXlsxRequest(BaseModel):
    rows: list[HierarchyExportRow]


@app.post("/api/hierarchy-analysis/export-xlsx")
def hierarchy_analysis_export_xlsx(body: HierarchyExportXlsxRequest, user: str = Depends(require_login)):
    """Excel export for the pending-primaries table - same approach as
    date_anomaly_detect_all_export_xlsx (see that route's docstring for
    why this round-trips through the server instead of a client-side
    Blob)."""
    if not body.rows:
        raise HTTPException(status_code=400, detail="No rows to export.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Hierarchy Analysis"
    headers = [
        "ID Measuring Point", "ID Main MP", "Supply (NISS)", "MP Type", "MP Status",
        "Secondaries", "Sec. Not Sent", "ID Calculation Module", "Type (Calc Module)",
        "Billing Period", "Billing Period Desc", "ID Reading", "Reading Date",
        "Reading Type", "Reading Type Desc", "Read Status", "Ready Usage", "Usage",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    # Column 17 (1-indexed) is "Ready Usage" - the analyst's own
    # most-important field on this page, bolded in the data rows too
    # (not just the header) so it stands out in the exported file the
    # same way it does in the on-screen table.
    ready_usage_col = 17
    for r in body.rows:
        ws.append([
            r.id_measuring_point, r.id_main_mp, r.niss, r.mp_type, r.mp_status,
            r.secondary_count, r.secondaries_not_sent_count, r.id_calculation_module, r.calc_module_type,
            r.id_billing_period, r.billing_period_desc, r.id_reading, r.reading_date,
            r.reading_type, r.reading_type_desc, r.read_status, r.ready_usage, r.reading_usage,
        ])
        ws.cell(row=ws.max_row, column=ready_usage_col).font = Font(bold=True)
    widths = [16, 14, 16, 12, 12, 12, 12, 16, 24, 14, 20, 14, 16, 14, 20, 14, 12, 10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="hierarchy_analysis.xlsx"'},
    )


# ---------------------------------------------------------------------
# Bulk Checker - ported from the standalone EWA Bulk Checker project
# (2026-09-12, RJ's request - see app/core/bulk_checker.py's module
# docstring for the full provenance note). Finds bulk accounts with no
# generated lot/file yet for a billing cycle and drills into one
# account's bills. Read-only against SQL Server (require_login only,
# same as every other read-heavy page here); notes/saved searches are
# local collaboration state any signed-in user can write, same as the
# source project's own "not admin-only" design for those.
# ---------------------------------------------------------------------
def _bc_parse_date(label: str, raw: str) -> date:
    try:
        return date.fromisoformat((raw or "").strip())
    except (ValueError, AttributeError):
        raise HTTPException(status_code=400, detail=f"{label} must be a valid date (YYYY-MM-DD).")


def _bc_parse_billing_period(raw) -> int:
    try:
        return int(str(raw).strip())
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="Billing period must be a number.")


@app.get("/api/bulk-checker/billing-periods")
def bulk_checker_billing_periods(user: str = Depends(require_login)):
    """Best-effort recent-periods picker - hidden client-side on failure,
    same as build_pending_bulks_sql's other table/column guesses (see
    bulk_checker.RECENT_BILLING_PERIODS_SQL's own comment). Typing a
    billing period by hand always works regardless."""
    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        return {"available": False, "columns": [], "rows": []}
    try:
        result = mssql.run_query(conn, bulk_checker.RECENT_BILLING_PERIODS_SQL)
    except mssql.ConnectionError_:
        return {"available": False, "columns": [], "rows": []}
    return {
        "available": True,
        "columns": result.columns,
        "rows": [[diff_engine.cell_display(v) for v in row] for row in result.rows],
    }


class BulkCheckerSearchRequest(BaseModel):
    date_from: str
    date_to: str
    billing_period: str
    status_filter: str = "all"


@app.post("/api/bulk-checker/search")
def bulk_checker_search(body: BulkCheckerSearchRequest, user: str = Depends(require_login)):
    date_from = _bc_parse_date("From date", body.date_from)
    date_to = _bc_parse_date("To date", body.date_to)
    if date_to <= date_from:
        raise HTTPException(status_code=400, detail="To date must be after the from date.")
    billing_period = _bc_parse_billing_period(body.billing_period)
    status_filter = body.status_filter if body.status_filter in bulk_checker.STATUS_FILTERS else "all"

    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    try:
        result = mssql.run_query(
            conn, bulk_checker.build_pending_bulks_sql(billing_period, date_from, date_to, status_filter)
        )
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    file_number_idx = next((i for i, c in enumerate(result.columns) if c.lower() == "file_number"), None)
    missing_bill_idx = next((i for i, c in enumerate(result.columns) if c.lower() == "has_missing_bill"), None)
    in_invoicing_idx = next((i for i, c in enumerate(result.columns) if c.lower() == "has_bill_in_invoicing"), None)
    pending_amount_idx = next((i for i, c in enumerate(result.columns) if c.lower() == "pending_amount"), None)

    columns = result.columns + ["is_pending"]
    display_rows: list[list[str]] = []
    pending_count = 0
    missing_bill_count = 0
    in_invoicing_count = 0
    outstanding_total = 0.0
    for row in result.rows:
        is_pending = file_number_idx is not None and row[file_number_idx] is None
        if is_pending:
            pending_count += 1
        if missing_bill_idx is not None and row[missing_bill_idx]:
            missing_bill_count += 1
        if in_invoicing_idx is not None and row[in_invoicing_idx]:
            in_invoicing_count += 1
        if (
            pending_amount_idx is not None
            # SQL Server money/decimal columns come back from pytds as
            # decimal.Decimal, not a plain float - the original isinstance
            # check here only covered (int, float) and silently treated
            # every real pending_amount as "not numeric", so Outstanding
            # always summed to 0 (caught live-verifying against the tunnel
            # DB, 2026-09-12: 242 rows with real non-zero pending_amount
            # values still showed "Outstanding: 0").
            and not is_pending
            and isinstance(row[pending_amount_idx], (int, float, decimal.Decimal))
            and row[pending_amount_idx] > 0
        ):
            outstanding_total += float(row[pending_amount_idx])
        display_rows.append([diff_engine.cell_display(v) for v in row] + ["Yes" if is_pending else "No"])

    # Snapshot into Trend history - only for a full, unfiltered search
    # (see bulk_checker_db.record_search's docstring for why a narrower
    # filter's row count isn't the period's true total). Best-effort:
    # never breaks the search itself.
    if status_filter == "all":
        try:
            bulk_checker_db.record_search(
                config.internal_db_path,
                billing_period=str(billing_period),
                date_from=date_from.isoformat(),
                date_to=date_to.isoformat(),
                total=len(result.rows),
                pending=pending_count,
                missing_bill=missing_bill_count,
                in_invoicing=in_invoicing_count,
                outstanding_amount=outstanding_total,
                searched_by=user,
            )
        except Exception:
            pass

    return {
        "columns": columns,
        "display_rows": display_rows,
        "row_count": len(result.rows),
        "pending_count": pending_count,
        "missing_bill_count": missing_bill_count,
        "in_invoicing_count": in_invoicing_count,
        "outstanding_amount": outstanding_total,
        "elapsed_ms": result.elapsed_ms,
        "status_filter": status_filter,
    }


class BulkCheckerDetailRequest(BaseModel):
    date_from: str
    date_to: str
    billing_period: str
    account_number: str
    bill_filter: str = "all"


@app.post("/api/bulk-checker/detail")
def bulk_checker_detail(body: BulkCheckerDetailRequest, user: str = Depends(require_login)):
    date_from = _bc_parse_date("From date", body.date_from)
    date_to = _bc_parse_date("To date", body.date_to)
    billing_period = _bc_parse_billing_period(body.billing_period)
    account_number = body.account_number.strip()
    if not account_number:
        raise HTTPException(status_code=400, detail="Account number is required.")
    bill_filter = body.bill_filter if body.bill_filter in bulk_checker.BILL_FILTERS else "all"

    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    try:
        result = mssql.run_query(
            conn, bulk_checker.build_bill_detail_sql(billing_period, date_from, date_to, account_number, bill_filter)
        )
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    return {
        "columns": result.columns,
        "display_rows": [[diff_engine.cell_display(v) for v in row] for row in result.rows],
        "row_count": result.row_count,
        "elapsed_ms": result.elapsed_ms,
        "bill_filter": bill_filter,
        "account_number": account_number,
    }


class BulkCheckerBulkDetailRequest(BaseModel):
    date_from: str
    date_to: str
    billing_period: str
    account_numbers: list[str]
    bill_filter: str = "all"


@app.post("/api/bulk-checker/bulk-detail")
def bulk_checker_bulk_detail(body: BulkCheckerBulkDetailRequest, user: str = Depends(require_login)):
    """Runs the same bill-detail query once per selected account and
    concatenates the results, so several accounts can be exported
    together instead of one at a time (same rationale as EWA's own
    pending_bulk_detail route)."""
    date_from = _bc_parse_date("From date", body.date_from)
    date_to = _bc_parse_date("To date", body.date_to)
    billing_period = _bc_parse_billing_period(body.billing_period)
    account_numbers = list(dict.fromkeys(a.strip() for a in body.account_numbers if a.strip()))
    if not account_numbers:
        raise HTTPException(status_code=400, detail="Select at least one account.")
    if len(account_numbers) > bulk_checker.MAX_BULK_ACCOUNTS:
        raise HTTPException(
            status_code=400, detail=f"Select at most {bulk_checker.MAX_BULK_ACCOUNTS} accounts at a time."
        )
    bill_filter = body.bill_filter if body.bill_filter in bulk_checker.BILL_FILTERS else "all"

    config = load_config()
    conn = config.get_active_connection()
    if not conn:
        raise HTTPException(status_code=400, detail="No connection configured.")
    columns: list[str] = []
    display_rows: list[list[str]] = []
    total_elapsed_ms = 0.0
    try:
        for account_number in account_numbers:
            result = mssql.run_query(
                conn, bulk_checker.build_bill_detail_sql(billing_period, date_from, date_to, account_number, bill_filter)
            )
            if not columns:
                columns = result.columns
            display_rows.extend([diff_engine.cell_display(v) for v in row] for row in result.rows)
            total_elapsed_ms += result.elapsed_ms
    except mssql.ConnectionError_ as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    return {
        "columns": columns,
        "display_rows": display_rows,
        "row_count": len(display_rows),
        "elapsed_ms": total_elapsed_ms,
        "account_numbers": account_numbers,
        "bill_filter": bill_filter,
    }


class BulkCheckerSaveSearchRequest(BaseModel):
    name: str
    date_from: str
    date_to: str
    billing_period: str


@app.get("/api/bulk-checker/saved-searches")
def bulk_checker_list_saved_searches(user: str = Depends(require_login)):
    config = load_config()
    return {
        "searches": [
            {
                "id": s.id, "name": s.name, "date_from": s.date_from, "date_to": s.date_to,
                "billing_period": s.billing_period, "created_by": s.created_by, "created_at_utc": s.created_at_utc,
            }
            for s in bulk_checker_db.list_saved_searches(config.internal_db_path)
        ]
    }


@app.post("/api/bulk-checker/saved-searches")
def bulk_checker_create_saved_search(body: BulkCheckerSaveSearchRequest, user: str = Depends(require_login)):
    _bc_parse_date("From date", body.date_from)
    _bc_parse_date("To date", body.date_to)
    _bc_parse_billing_period(body.billing_period)
    config = load_config()
    try:
        s = bulk_checker_db.save_search(
            config.internal_db_path, body.name, body.date_from, body.date_to, body.billing_period, user
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {
        "id": s.id, "name": s.name, "date_from": s.date_from, "date_to": s.date_to,
        "billing_period": s.billing_period,
    }


@app.delete("/api/bulk-checker/saved-searches/{search_id}")
def bulk_checker_delete_saved_search(search_id: int, user: str = Depends(require_login)):
    config = load_config()
    bulk_checker_db.delete_saved_search(config.internal_db_path, search_id)
    return {"ok": True}


@app.get("/api/bulk-checker/trend")
def bulk_checker_trend(user: str = Depends(require_login)):
    config = load_config()
    entries = bulk_checker_db.list_search_history(config.internal_db_path)
    return {
        "entries": [
            {
                "billing_period": e.billing_period, "date_from": e.date_from, "date_to": e.date_to,
                "total": e.total, "pending": e.pending, "missing_bill": e.missing_bill,
                "in_invoicing": e.in_invoicing, "outstanding_amount": e.outstanding_amount,
                "searched_by": e.searched_by, "searched_at_utc": e.searched_at_utc,
            }
            for e in entries
        ]
    }


class BulkCheckerNoteRequest(BaseModel):
    billing_period: str
    status: str = bulk_checker_db.STATUS_OPEN
    note: str = ""


@app.get("/api/bulk-checker/notes")
def bulk_checker_list_notes(billing_period: str, user: str = Depends(require_login)):
    config = load_config()
    items = bulk_checker_db.list_notes_for_period(config.internal_db_path, billing_period)
    return {
        "notes": {
            n.account_number: {
                "status": n.status, "note": n.note,
                "updated_by": n.updated_by, "updated_at_utc": n.updated_at_utc,
            }
            for n in items
        }
    }


@app.put("/api/bulk-checker/notes/{account_number}")
def bulk_checker_upsert_note(account_number: str, body: BulkCheckerNoteRequest, user: str = Depends(require_login)):
    config = load_config()
    try:
        n = bulk_checker_db.set_note(
            config.internal_db_path, account_number, body.billing_period, body.status, body.note, user
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {
        "account_number": n.account_number, "status": n.status, "note": n.note,
        "updated_by": n.updated_by, "updated_at_utc": n.updated_at_utc,
    }


@app.delete("/api/bulk-checker/notes/{account_number}")
def bulk_checker_delete_note(account_number: str, billing_period: str, user: str = Depends(require_login)):
    config = load_config()
    bulk_checker_db.clear_note(config.internal_db_path, account_number, billing_period)
    return {"ok": True}


class BulkCheckerExportXlsxRequest(BaseModel):
    """Generic (not per-field) export request - Bulk Checker's two result
    sets are wide (13 / 29 columns) and come straight from SQL Server
    column names, unlike Hierarchy's small curated field set, so the
    frontend just sends back whatever headers/rows are currently
    visible/filtered on screen rather than this needing a fixed-shape
    row model per column."""

    filename: str = "bulk_checker.xlsx"
    headers: list[str]
    rows: list[list[str]]


@app.post("/api/bulk-checker/export-xlsx")
def bulk_checker_export_xlsx(body: BulkCheckerExportXlsxRequest, user: str = Depends(require_login)):
    if not body.rows:
        raise HTTPException(status_code=400, detail="No rows to export.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Bulk Checker"
    ws.append(body.headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in body.rows:
        ws.append(row)
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    filename = body.filename if body.filename.endswith(".xlsx") else f"{body.filename}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------
# Static frontend (must be mounted last so /api/* routes above win)
# ---------------------------------------------------------------------
app.mount("/", StaticFiles(directory=str(STATIC_DIR), html=True), name="static")
