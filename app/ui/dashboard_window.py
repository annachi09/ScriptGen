"""
Dashboard window: a separate Toplevel showing summary statistics and
charts for the current query result. Opened on demand from the main
window - never blocks it, and reads a snapshot of the grid at the
moment it was opened (it does not stay live-synced to further edits).

Three tabs:
  - Overview: KPI cards, a searchable per-column stats table, and two
    charts (numeric distribution / category breakdown), each exportable.
  - Trend History: if this result's source table has 2+ prior exports
    in the internal SQLite snapshot store, plots row count over time
    for that table - a lightweight "is this table growing/shrinking"
    view that costs nothing extra since the snapshot data already
    exists (app/db/internal_store.py).
  - Correlation: pick two numeric columns and see a scatter plot plus
    their Pearson r (app.core.stats.pearson_correlation) - whether two
    numbers in this result move together, e.g. "does amount correlate
    with days_open".

The whole dashboard (every chart that's currently drawn) can also be
exported as one multi-page PDF via the "Export Dashboard PDF" button
in the summary bar.
"""
from __future__ import annotations

import tkinter as tk
from datetime import datetime, timezone

import ttkbootstrap as tb
from ttkbootstrap.constants import *
from tkinter import filedialog

from ..core import stats as stats_mod
from ..db import internal_store
from ..db.mssql import QueryResult
from .theme import PAD, PAD_SM, is_dark

_MPL_DARK_BG = "#282c34"
_MPL_DARK_FG = "#d7dae0"
# One consistent hue per chart "job" (dataviz skill: sequential = one hue,
# categorical identity gets its own distinct hue) rather than a different
# color per bar/card - matches how the rest of the app already colors
# things (blue = numeric/quantitative, amber = categorical/breakdown).
_NUMERIC_HUE = "#3b82f6"
_CATEGORICAL_HUE = "#f59e0b"
_TREND_HUE = "#3b82f6"
_CORRELATION_HUE = "#8b5cf6"

# matplotlib is a genuinely slow import (~0.5-2s, more in a frozen exe) -
# importing it at module level here means EVERY app launch pays that cost
# before the main window can even appear, whether or not the Dashboard is
# ever opened, since app/ui/main_window.py imports this module up front.
# Deferring it to first-use (here, the first time a DashboardWindow is
# actually constructed) keeps startup fast for the common case and only
# costs the import once per run, the first time it's genuinely needed.
Figure = None
FigureCanvasTkAgg = None


def _ensure_matplotlib():
    global Figure, FigureCanvasTkAgg
    if Figure is None:
        from matplotlib.figure import Figure as _Figure
        from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg as _FigureCanvasTkAgg

        Figure = _Figure
        FigureCanvasTkAgg = _FigureCanvasTkAgg


class DashboardWindow(tb.Toplevel):
    def __init__(self, parent, query_result: QueryResult, theme: str, internal_db_path: str = ""):
        _ensure_matplotlib()
        super().__init__(parent)
        title_table = query_result.source_table or "query result"
        self.title(f"Dashboard - {title_table}")
        self.geometry("1120x800")
        self.minsize(860, 600)

        self.query_result = query_result
        self.dark = is_dark(theme)
        self.internal_db_path = internal_db_path
        self.stats = stats_mod.compute_column_stats(query_result.columns, query_result.rows)
        self.numeric_cols = stats_mod.numeric_columns(self.stats)
        self.categorical_cols = stats_mod.categorical_columns(self.stats)
        self.summary = stats_mod.summary_counts(self.stats)

        self._build_summary_bar()

        notebook = tb.Notebook(self)
        notebook.pack(fill="both", expand=True, padx=PAD, pady=(0, PAD))
        overview = tb.Frame(notebook, padding=(0, PAD_SM, 0, 0))
        history = tb.Frame(notebook, padding=PAD)
        correlation = tb.Frame(notebook, padding=PAD)
        notebook.add(overview, text="  Overview  ")
        notebook.add(history, text="  Trend History  ")
        notebook.add(correlation, text="  Correlation  ")

        self._build_kpi_cards(overview)
        self._build_stats_table(overview)
        self._build_charts(overview)
        self._build_trend_history(history)
        self._build_correlation_tab(correlation)

    # ------------------------------------------------------------------
    def _build_summary_bar(self):
        bar = tb.Frame(self, bootstyle=PRIMARY, padding=(PAD, PAD_SM))
        bar.pack(fill="x")

        top_row = tb.Frame(bar, bootstyle=PRIMARY)
        top_row.pack(fill="x")
        tb.Label(
            top_row, text="Query Result Dashboard", font=("TkDefaultFont", 13, "bold"), bootstyle="inverse-primary"
        ).pack(side="left", anchor="w")
        tb.Button(
            top_row, text="⬇ Export Dashboard PDF...", bootstyle="light-outline",
            command=self._export_dashboard_pdf,
        ).pack(side="right")

        generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        summary = f"query took {self.query_result.elapsed_ms:.0f} ms   ·   generated {generated}"
        tb.Label(bar, text=summary, font=("TkDefaultFont", 9), bootstyle="inverse-primary").pack(anchor="w")
        if self.query_result.source_table:
            tb.Label(
                bar, text=f"Source table (best guess): {self.query_result.source_table}",
                font=("TkDefaultFont", 9), bootstyle="inverse-primary",
            ).pack(anchor="w")

    # ------------------------------------------------------------------
    # KPI cards - one glance for "how big / how clean is this result"
    # before drilling into the per-column table. Deliberately one flat
    # card style (not a different color per card) - color here would be
    # decorative, not identity, which is the anti-pattern the dataviz
    # skill calls out; a matching border/label style for every card keeps
    # the row reading as one system.
    # ------------------------------------------------------------------
    def _build_kpi_cards(self, parent):
        row = tb.Frame(parent)
        row.pack(fill="x", padx=PAD, pady=(0, PAD_SM))

        cards = [
            ("Rows", f"{self.query_result.row_count:,}"),
            ("Columns", str(len(self.query_result.columns))),
            ("Numeric cols", str(self.summary["numeric_columns"])),
            ("Text cols", str(self.summary["categorical_columns"])),
            ("Empty cols", str(self.summary["empty_columns"])),
            ("Total NULLs", f"{self.summary['total_nulls']:,}"),
        ]
        for i, (label, value) in enumerate(cards):
            row.grid_columnconfigure(i, weight=1)
            card = tb.Frame(row, bootstyle="secondary", padding=PAD_SM)
            card.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else PAD_SM, 0))
            tb.Label(card, text=value, font=("TkDefaultFont", 15, "bold"), bootstyle="inverse-secondary").pack(anchor="w")
            tb.Label(card, text=label, font=("TkDefaultFont", 8), bootstyle="inverse-secondary").pack(anchor="w")

    # ------------------------------------------------------------------
    def _build_stats_table(self, parent):
        card = tb.Labelframe(parent, text="Per-Column Summary", padding=PAD_SM, bootstyle="secondary")
        card.pack(fill="x", padx=PAD, pady=(0, PAD_SM))

        toolbar = tb.Frame(card)
        toolbar.pack(fill="x", pady=(0, PAD_SM))
        tb.Label(toolbar, text="Filter:", bootstyle="secondary").pack(side="left")
        self.stats_filter_var = tk.StringVar()
        filter_entry = tb.Entry(toolbar, textvariable=self.stats_filter_var, width=24)
        filter_entry.pack(side="left", padx=(PAD_SM, 0))
        filter_entry.bind("<KeyRelease>", lambda e: self._apply_stats_filter())
        export_btn = tb.Menubutton(toolbar, text="⬇ Export Stats", bootstyle="secondary-outline", direction="below")
        export_btn.pack(side="right")
        export_menu = tk.Menu(export_btn, tearoff=0)
        export_btn["menu"] = export_menu
        export_menu.add_command(label="CSV...", command=self._export_stats_csv)
        export_menu.add_command(label="Excel (.xlsx)...", command=self._export_stats_excel)

        columns = ("type", "non_null", "nulls", "distinct", "min", "max", "mean", "top_value")
        self.stats_tree = tb.Treeview(card, columns=columns, show="tree headings", height=min(8, len(self.stats) + 1))
        self.stats_tree.heading("#0", text="Column")
        self.stats_tree.column("#0", width=140, anchor="w")
        headings = {
            "type": "Type", "non_null": "Non-null", "nulls": "Nulls", "distinct": "Distinct",
            "min": "Min", "max": "Max", "mean": "Mean", "top_value": "Top value (count)",
        }
        for col in columns:
            self.stats_tree.heading(col, text=headings[col])
            self.stats_tree.column(col, width=100, anchor="center")
        self.stats_tree.column("top_value", width=160, anchor="w")

        self._stats_rows: dict[str, tuple] = {}
        for name, s in self.stats.items():
            if s.kind == "numeric":
                row = (
                    "numeric", s.non_null_count, s.null_count, s.distinct_count,
                    f"{s.min_value:g}", f"{s.max_value:g}", f"{s.mean_value:.2f}", "",
                )
            elif s.kind == "categorical":
                top = f"{s.top_values[0][0]} ({s.top_values[0][1]})" if s.top_values else ""
                row = ("text", s.non_null_count, s.null_count, s.distinct_count, "", "", "", top)
            else:
                row = ("empty", 0, s.null_count, 0, "", "", "", "")
            self._stats_rows[name] = row

        self._apply_stats_filter()
        self.stats_tree.pack(fill="x")

    def _apply_stats_filter(self):
        needle = self.stats_filter_var.get().strip().lower()
        self.stats_tree.delete(*self.stats_tree.get_children())
        for name, row in self._stats_rows.items():
            if needle and needle not in name.lower():
                continue
            self.stats_tree.insert("", "end", text=name, values=row)

    def _export_stats_csv(self):
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV file", "*.csv")])
        if not path:
            return
        import csv

        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["column", "type", "non_null", "nulls", "distinct", "min", "max", "mean", "top_value"])
            for name, row in self._stats_rows.items():
                writer.writerow([name, *row])

    def _export_stats_excel(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        from openpyxl import Workbook
        from openpyxl.styles import Font

        header = ["column", "type", "non_null", "nulls", "distinct", "min", "max", "mean", "top_value"]
        wb = Workbook()
        ws = wb.active
        ws.title = "Per-Column Summary"
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for name, row in self._stats_rows.items():
            ws.append([name, *row])
        for col_idx, col_name in enumerate(header, start=1):
            longest = len(col_name)
            for name, row in self._stats_rows.items():
                values = [name, *row]
                longest = max(longest, len(str(values[col_idx - 1])))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(60, max(10, longest + 2))
        wb.save(path)

    # ------------------------------------------------------------------
    def _build_charts(self, parent):
        charts_frame = tb.Frame(parent)
        charts_frame.pack(fill="both", expand=True, padx=PAD, pady=(0, PAD))

        left = tb.Labelframe(charts_frame, text="Numeric Distribution", padding=PAD_SM, bootstyle="info")
        left.pack(side="left", fill="both", expand=True, padx=(0, PAD_SM))

        right = tb.Labelframe(charts_frame, text="Category Breakdown", padding=PAD_SM, bootstyle="warning")
        right.pack(side="left", fill="both", expand=True, padx=(PAD_SM, 0))

        self._build_numeric_panel(left)
        self._build_categorical_panel(right)

    def _style_axes(self, fig: Figure, ax):
        if self.dark:
            fig.patch.set_facecolor(_MPL_DARK_BG)
            ax.set_facecolor(_MPL_DARK_BG)
            ax.tick_params(colors=_MPL_DARK_FG, labelsize=8)
            for spine in ax.spines.values():
                spine.set_color(_MPL_DARK_FG)
            ax.title.set_color(_MPL_DARK_FG)
            ax.xaxis.label.set_color(_MPL_DARK_FG)
            ax.yaxis.label.set_color(_MPL_DARK_FG)
        else:
            ax.tick_params(labelsize=8)

    def _save_figure(self, fig: Figure):
        path = filedialog.asksaveasfilename(defaultextension=".png", filetypes=[("PNG image", "*.png")])
        if not path:
            return
        fig.savefig(path, facecolor=fig.get_facecolor())

    def _build_numeric_panel(self, parent):
        if not self.numeric_cols:
            tb.Label(parent, text="No numeric columns in this result.", bootstyle="secondary").pack(
                expand=True, pady=40
            )
            return

        control_row = tb.Frame(parent)
        control_row.pack(fill="x", pady=(0, PAD_SM))
        tb.Label(control_row, text="Column:").pack(side="left")
        self.numeric_choice = tk.StringVar(value=self.numeric_cols[0])
        combo = tb.Combobox(control_row, textvariable=self.numeric_choice, values=self.numeric_cols, state="readonly", width=16)
        combo.pack(side="left", padx=PAD_SM)
        combo.bind("<<ComboboxSelected>>", lambda e: self._redraw_numeric_chart())
        tb.Button(
            control_row, text="💾", bootstyle="secondary-outline", width=3,
            command=lambda: self._save_figure(self.numeric_fig),
        ).pack(side="right")

        self.numeric_fig = Figure(figsize=(4.6, 3.4), dpi=100)
        self.numeric_canvas = FigureCanvasTkAgg(self.numeric_fig, master=parent)
        self.numeric_canvas.get_tk_widget().pack(fill="both", expand=True)
        self._redraw_numeric_chart()

    def _redraw_numeric_chart(self):
        col = self.numeric_choice.get()
        col_index = self.query_result.columns.index(col)
        labels, counts = stats_mod.numeric_bucket_histogram(self.query_result.rows, col_index)

        self.numeric_fig.clear()
        ax = self.numeric_fig.add_subplot(111)
        if labels:
            ax.bar(range(len(labels)), counts, color=_NUMERIC_HUE)
            ax.set_xticks(range(len(labels)))
            ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=7)
        ax.set_title(col, fontsize=10)
        ax.set_ylabel("count", fontsize=8)
        self._style_axes(self.numeric_fig, ax)
        self.numeric_fig.tight_layout()
        self.numeric_canvas.draw()

    def _build_categorical_panel(self, parent):
        if not self.categorical_cols:
            tb.Label(parent, text="No categorical/text columns in this result.", bootstyle="secondary").pack(
                expand=True, pady=40
            )
            return

        control_row = tb.Frame(parent)
        control_row.pack(fill="x", pady=(0, PAD_SM))
        tb.Label(control_row, text="Column:").pack(side="left")
        self.cat_choice = tk.StringVar(value=self.categorical_cols[0])
        combo = tb.Combobox(control_row, textvariable=self.cat_choice, values=self.categorical_cols, state="readonly", width=16)
        combo.pack(side="left", padx=PAD_SM)
        combo.bind("<<ComboboxSelected>>", lambda e: self._redraw_categorical_chart())
        tb.Button(
            control_row, text="💾", bootstyle="secondary-outline", width=3,
            command=lambda: self._save_figure(self.cat_fig),
        ).pack(side="right")

        self.cat_fig = Figure(figsize=(4.6, 3.4), dpi=100)
        self.cat_canvas = FigureCanvasTkAgg(self.cat_fig, master=parent)
        self.cat_canvas.get_tk_widget().pack(fill="both", expand=True)
        self._redraw_categorical_chart()

    def _redraw_categorical_chart(self):
        col = self.cat_choice.get()
        top_values = self.stats[col].top_values[:8]

        self.cat_fig.clear()
        ax = self.cat_fig.add_subplot(111)
        if top_values:
            labels = [v for v, _ in reversed(top_values)]
            counts = [c for _, c in reversed(top_values)]
            ax.barh(range(len(labels)), counts, color=_CATEGORICAL_HUE)
            ax.set_yticks(range(len(labels)))
            ax.set_yticklabels(labels, fontsize=7)
        ax.set_title(col, fontsize=10)
        ax.set_xlabel("count", fontsize=8)
        self._style_axes(self.cat_fig, ax)
        self.cat_fig.tight_layout()
        self.cat_canvas.draw()

    # ------------------------------------------------------------------
    # Trend History - free insight from data that already exists: every
    # snapshot export (internal_store) for this same source table is a
    # data point of "row count at time T". Two or more of them and we can
    # draw a one-line trend; fewer than that, or no db path / source
    # table at all, and there's nothing honest to plot, so we say so
    # instead of drawing an empty or single-point "chart" (dataviz skill:
    # not every summary is a chart - a message/stat tile can be the
    # correct answer).
    # ------------------------------------------------------------------
    def _build_trend_history(self, parent):
        parent.grid_columnconfigure(0, weight=1)
        parent.grid_rowconfigure(1, weight=1)

        source_table = self.query_result.source_table
        if not source_table or not self.internal_db_path:
            self.trend_snapshot_count = 0
            tb.Label(
                parent,
                text=(
                    "No source table detected for this result (or no internal DB configured), "
                    "so there's no snapshot history to compare against.\nExport a few snapshots "
                    "of the same table over time (File > Export Result to Internal DB) to see a "
                    "row-count trend here."
                ),
                bootstyle="secondary", wraplength=760, justify="left",
            ).grid(row=0, column=0, sticky="w", pady=40)
            return

        try:
            all_snapshots = internal_store.list_snapshots(self.internal_db_path)
        except Exception as exc:  # local sqlite file issues shouldn't crash the dashboard
            self.trend_snapshot_count = 0
            tb.Label(parent, text=f"Could not read snapshot history: {exc}", bootstyle="danger").grid(
                row=0, column=0, sticky="w"
            )
            return

        matching = [s for s in all_snapshots if s.source_table == source_table]
        matching.sort(key=lambda s: s.created_at_utc)
        self.trend_snapshot_count = len(matching)  # exposed for tests; the widgets below are the real UI

        header = tb.Label(
            parent, text=f"Row count over time for {source_table}", font=("TkDefaultFont", 11, "bold")
        )
        header.grid(row=0, column=0, sticky="w", pady=(0, PAD_SM))

        if len(matching) < 2:
            tb.Label(
                parent,
                text=(
                    f"Only {len(matching)} snapshot(s) exported for {source_table} so far - need at least 2 "
                    "to draw a trend. Export another snapshot later (File > Export Result to Internal DB) "
                    "and reopen the Dashboard to see it here."
                ),
                bootstyle="secondary", wraplength=760, justify="left",
            ).grid(row=1, column=0, sticky="nw")
            return

        body = tb.Frame(parent)
        body.grid(row=1, column=0, sticky="nsew")
        body.grid_columnconfigure(0, weight=3)
        body.grid_columnconfigure(1, weight=2)
        body.grid_rowconfigure(0, weight=1)

        chart_frame = tb.Frame(body)
        chart_frame.grid(row=0, column=0, sticky="nsew", padx=(0, PAD))

        fig = Figure(figsize=(5.2, 3.6), dpi=100)
        self.trend_fig = fig  # exposed for _export_dashboard_pdf
        canvas = FigureCanvasTkAgg(fig, master=chart_frame)
        ax = fig.add_subplot(111)

        x = list(range(len(matching)))
        y = [s.row_count for s in matching]
        labels = [s.created_at_utc[:16].replace("T", " ") for s in matching]
        ax.plot(x, y, color=_TREND_HUE, linewidth=2, marker="o", markersize=6)
        ax.set_xticks(x)
        ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=7)
        ax.set_ylabel("row count", fontsize=8)
        ax.set_title(f"{len(matching)} snapshots", fontsize=10)
        self._style_axes(fig, ax)
        fig.tight_layout()
        canvas.get_tk_widget().pack(fill="both", expand=True)
        canvas.draw()

        tb.Button(
            chart_frame, text="💾 Save Chart PNG", bootstyle="secondary-outline",
            command=lambda: self._save_figure(fig),
        ).pack(anchor="e", pady=(PAD_SM, 0))

        list_frame = tb.Labelframe(body, text="Snapshots", padding=PAD_SM, bootstyle="secondary")
        list_frame.grid(row=0, column=1, sticky="nsew")
        tree = tb.Treeview(list_frame, columns=("rows",), show="tree headings")
        tree.heading("#0", text="Exported (UTC)")
        tree.heading("rows", text="Rows")
        tree.column("#0", width=150)
        tree.column("rows", width=70, anchor="e")
        for s in matching:
            tree.insert("", "end", text=s.created_at_utc[:16].replace("T", " "), values=(s.row_count,))
        tree.pack(fill="both", expand=True)

    # ------------------------------------------------------------------
    # Correlation - scatter + Pearson r between two numeric columns.
    # Genuinely needs 2+ numeric columns to mean anything, so - same
    # philosophy as Trend History above - a clear message stands in for
    # the chart when that isn't the case, rather than an empty plot.
    # ------------------------------------------------------------------
    def _build_correlation_tab(self, parent):
        parent.grid_columnconfigure(0, weight=1)
        parent.grid_rowconfigure(1, weight=1)

        if len(self.numeric_cols) < 2:
            tb.Label(
                parent,
                text=(
                    "Need at least 2 numeric columns in this result to compute a correlation - "
                    f"this result has {len(self.numeric_cols)}."
                ),
                bootstyle="secondary", wraplength=760, justify="left",
            ).grid(row=0, column=0, sticky="w", pady=40)
            return

        control_row = tb.Frame(parent)
        control_row.grid(row=0, column=0, sticky="ew", pady=(0, PAD_SM))
        tb.Label(control_row, text="X:").pack(side="left")
        self.corr_x_choice = tk.StringVar(value=self.numeric_cols[0])
        combo_x = tb.Combobox(
            control_row, textvariable=self.corr_x_choice, values=self.numeric_cols, state="readonly", width=16
        )
        combo_x.pack(side="left", padx=(PAD_SM, PAD))
        tb.Label(control_row, text="Y:").pack(side="left")
        default_y = self.numeric_cols[1] if len(self.numeric_cols) > 1 else self.numeric_cols[0]
        self.corr_y_choice = tk.StringVar(value=default_y)
        combo_y = tb.Combobox(
            control_row, textvariable=self.corr_y_choice, values=self.numeric_cols, state="readonly", width=16
        )
        combo_y.pack(side="left", padx=(PAD_SM, 0))
        combo_x.bind("<<ComboboxSelected>>", lambda e: self._redraw_correlation_chart())
        combo_y.bind("<<ComboboxSelected>>", lambda e: self._redraw_correlation_chart())
        tb.Button(
            control_row, text="💾", bootstyle="secondary-outline", width=3,
            command=lambda: self._save_figure(self.corr_fig),
        ).pack(side="right")

        self.corr_stat_var = tk.StringVar()
        tb.Label(
            parent, textvariable=self.corr_stat_var, font=("TkDefaultFont", 11, "bold"), bootstyle="secondary"
        ).grid(row=0, column=0, sticky="e", padx=(0, 40))

        self.corr_fig = Figure(figsize=(6.0, 4.2), dpi=100)
        self.corr_canvas = FigureCanvasTkAgg(self.corr_fig, master=parent)
        self.corr_canvas.get_tk_widget().grid(row=1, column=0, sticky="nsew")
        self._redraw_correlation_chart()

    @staticmethod
    def _correlation_strength_label(r: float) -> str:
        magnitude = abs(r)
        if magnitude >= 0.7:
            strength = "strong"
        elif magnitude >= 0.4:
            strength = "moderate"
        elif magnitude >= 0.1:
            strength = "weak"
        else:
            strength = "negligible"
        direction = "positive" if r > 0 else "negative" if r < 0 else "no"
        return f"{strength} {direction}" if magnitude >= 0.1 else "negligible"

    def _redraw_correlation_chart(self):
        col_x = self.corr_x_choice.get()
        col_y = self.corr_y_choice.get()
        idx_x = self.query_result.columns.index(col_x)
        idx_y = self.query_result.columns.index(col_y)

        r = stats_mod.pearson_correlation(self.query_result.rows, idx_x, idx_y)
        if r is None:
            self.corr_stat_var.set("Pearson r = n/a (not enough overlapping numeric values)")
        else:
            self.corr_stat_var.set(f"Pearson r = {r:.3f}  ({self._correlation_strength_label(r)})")

        xs, ys = [], []
        for row in self.query_result.rows:
            x_val, y_val = row[idx_x], row[idx_y]
            try:
                xs.append(float(x_val))
                ys.append(float(y_val))
            except (TypeError, ValueError):
                continue

        self.corr_fig.clear()
        ax = self.corr_fig.add_subplot(111)
        if xs:
            ax.scatter(xs, ys, color=_CORRELATION_HUE, alpha=0.7, edgecolors="none")
        ax.set_xlabel(col_x, fontsize=8)
        ax.set_ylabel(col_y, fontsize=8)
        ax.set_title(f"{col_x} vs {col_y}", fontsize=10)
        self._style_axes(self.corr_fig, ax)
        self.corr_fig.tight_layout()
        self.corr_canvas.draw()

    # ------------------------------------------------------------------
    # Dashboard-wide PDF export - one page per chart that's currently on
    # screen (numeric distribution, category breakdown, trend, and
    # correlation, whichever of those actually got built for this
    # result), so someone can hand the whole picture to someone else
    # without screenshotting every tab.
    # ------------------------------------------------------------------
    def _export_dashboard_pdf(self):
        figs = []
        for attr, label in (
            ("numeric_fig", "Numeric Distribution"),
            ("cat_fig", "Category Breakdown"),
            ("trend_fig", "Trend History"),
            ("corr_fig", "Correlation"),
        ):
            fig = getattr(self, attr, None)
            if fig is not None:
                figs.append((fig, label))

        if not figs:
            from tkinter import messagebox

            messagebox.showinfo("Nothing to export", "No charts are available for this result yet.")
            return

        path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF document", "*.pdf")])
        if not path:
            return

        from matplotlib.backends.backend_pdf import PdfPages

        with PdfPages(path) as pdf:
            for fig, _label in figs:
                pdf.savefig(fig, facecolor=fig.get_facecolor())
