"""
Headless smoke test: launches the real MainWindow (ttkbootstrap themed)
inside Xvfb against a fake query result (no real SQL Server needed) and
exercises essentially every UI-level feature end to end - the
key-column picker, grid highlighting, script generation, CSV/Excel/
internal-DB export, the Dashboard window (KPI cards, stats
filter/export, Trend History, Correlation tab, dashboard-wide PDF
export), grid perf on a large result, SQL syntax highlighting, Cancel
Query, the mandatory Jira/program field, and every section of the
Tools tab (Saved Queries, Schema Validation, Rollback Script, Snapshot
Diff Viewer, Environment switcher), plus the AI Assist page's NL-WHERE
builder and the Script page's AI Review button (both the offline
"no key configured" short-circuit and a full round trip with the
Gemini call itself monkeypatched, so this stays fast and offline).
Catches import errors, Tk/ttkbootstrap widget construction errors, and
layout/wiring exceptions that the pure-logic unit tests can't see.

Threaded AI/DB-lookup actions normally hand their result back via
root.after(0, ...) from a background thread, which only works with a
genuine root.mainloop() running - this test drives Tk via repeated
root.update() instead, so for the one flow that needs a real round
trip (the AI actions) it swaps in a synchronous stand-in for
threading.Thread just for that block, restoring it immediately after.

Run with:  xvfb-run -a python tests/smoke_launch.py
"""
import random
import string
import sys
import time
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# Hermetic run: this test asserts against a LIGHT_THEME starting state and
# reads/writes the real data/config.json (same path a dev-mode run of the
# app itself would use - see app/utils/paths.py). A previous run of this
# same script (e.g. the theme-toggle assertion at the bottom, which
# persists DARK_THEME) would otherwise leak into this run's "starting"
# state. Wiped unconditionally before anything imports/uses config.py.
_data_dir = Path(__file__).resolve().parent.parent / "data"
for _stale in ("config.json", "scriptgen.key"):
    (_data_dir / _stale).unlink(missing_ok=True)

import ttkbootstrap as tb
from tkinter import messagebox

from app.ui.theme import LIGHT_THEME, DARK_THEME
from app.core import schema_check, ai_assist


def main():
    root = tb.Window(themename=LIGHT_THEME)
    try:
        from app.ui.main_window import MainWindow
        from app.db.mssql import QueryResult

        win = MainWindow(root)
        root.update_idletasks()
        root.update()

        assert win.sheet is not None
        assert win.config is not None

        # Simulate a query result to exercise the key-column picker path
        fake_result = QueryResult(
            columns=["id", "name", "amount"],
            rows=[[1, "Alice", 10], [2, "Bob", 20]],
            elapsed_ms=5.0,
            source_table="dbo.Accounts",
            source_schema="dbo",
        )
        win._on_query_success(fake_result)
        root.update_idletasks()
        root.update()
        assert set(win.key_column_vars.keys()) == {"id", "name", "amount"}

        win._select_all_keys()
        assert win._selected_key_columns() == ["id", "name", "amount"]
        win._clear_all_keys()
        assert win._selected_key_columns() == []

        # The WHERE-key picker / target fields / Generate button live in
        # the resizable sidebar now, not a bottom bar on the Workspace
        # page - check the sash exists, has a sane starting position, and
        # can actually be moved (drag-to-resize).
        assert hasattr(win, "body_paned"), "sidebar should be a resizable Panedwindow pane"
        sash = win.body_paned.sashpos(0)
        assert 100 < sash < 900, f"unexpected initial sash position: {sash}"
        win.body_paned.sashpos(0, sash + 80)
        root.update_idletasks()
        assert win.body_paned.sashpos(0) == sash + 80, "sidebar should be draggable/resizable"

        # Grid highlighting: editing a cell should highlight that cell and
        # its row's index, and generating a script should call out which
        # column(s) changed - both should clear back to nothing once the
        # sheet is repopulated with a fresh result.
        win._on_query_success(fake_result)
        root.update_idletasks()
        data = win.sheet.get_sheet_data()
        data[0][2] = "999"  # was 10
        win.sheet.set_sheet_data(data)
        win._refresh_grid_highlights()
        root.update_idletasks()

        cell_highlights = win.sheet.get_highlighted_cells(canvas="table")
        # tksheet's setter (highlight_cells) accepts "index" as an alias for
        # "row_index", but its getter (get_highlighted_cells) only recognizes
        # "row_index" - "index" silently returns None there. Use the getter's
        # real name here.
        row_highlights = win.sheet.get_highlighted_cells(canvas="row_index")
        assert cell_highlights, "editing a cell should highlight it"
        assert (0, 2) in cell_highlights, f"expected (row0, col2) highlighted, got {cell_highlights}"
        assert row_highlights, "the edited row's index should be highlighted too"
        assert 0 in row_highlights, f"expected row 0's index highlighted, got {row_highlights}"

        win.target_schema_var.set("dbo")
        win.target_table_var.set("Accounts")
        win.key_column_vars["id"].set(True)
        root.update_idletasks()

        # SQL Preview companion sheet: the edited row (row 0, amount ->
        # 999) should now show its own live UPDATE statement right next
        # to it, row-synced with the main grid; the untouched row (row
        # 1) should stay blank. Setting Target schema/table/key above
        # should have been enough on its own to populate this (via the
        # trace_add hooks), with no explicit _refresh_grid_highlights()
        # call needed here - that's the point of wiring it that way.
        assert hasattr(win, "sql_preview_sheet"), "SQL Preview companion sheet should exist"
        preview_data = win.sql_preview_sheet.get_sheet_data()
        assert preview_data[0][0], "edited row should have a non-blank SQL preview"
        assert "UPDATE dbo.Accounts" in preview_data[0][0], preview_data[0][0]
        # amount's original value was an int (10), so the coerced edit
        # round-trips back to an int (999) and renders as a bare number,
        # not a quoted string - see diff_engine.coerce_edited_value.
        assert "amount = 999" in preview_data[0][0], preview_data[0][0]
        assert "WHERE id = 1" in preview_data[0][0], preview_data[0][0]
        # The WHERE clause should also guard on amount's ORIGINAL value
        # (10), not just the id key - so the statement only fires if the
        # row still holds the value the grid was generated from.
        assert "WHERE id = 1 AND amount = 10" in preview_data[0][0], preview_data[0][0]
        assert preview_data[1][0] == "", f"unedited row should have a blank preview, got {preview_data[1][0]!r}"
        # No brackets anywhere in the preview - same unquoting as the full script.
        assert "[" not in preview_data[0][0] and "]" not in preview_data[0][0], preview_data[0][0]

        win._generate_script()
        script_text = win.script_text.text.get("1.0", "end")
        assert "changed column(s): amount" in script_text, script_text

        # Re-running the query should reset to a clean, unhighlighted grid
        # and blank out every row's SQL preview (nothing edited yet).
        win._on_query_success(fake_result)
        root.update_idletasks()
        assert not win.sheet.get_highlighted_cells(canvas="table"), "fresh query result should start unhighlighted"
        assert all(row[0] == "" for row in win.sql_preview_sheet.get_sheet_data()), (
            "a freshly (re)loaded result should start with every SQL preview blank"
        )

        # Inline Gemini key setup on the AI Assist page: Save should
        # persist to config (encrypted at rest, decrypted back out here)
        # and flip the status line to the "configured" state; Explain
        # should exist as a callable action alongside Suggest/Optimize.
        assert hasattr(win, "ai_key_var") and hasattr(win, "ai_key_status_var")
        win.ai_key_var.set("fake-test-key-1234567890")
        win._save_ai_key()
        assert win.config.ai.get_api_key() == "fake-test-key-1234567890"
        assert win.config.ai.enabled is True
        assert win.ai_key_status_var.get().startswith("✓"), win.ai_key_status_var.get()
        assert hasattr(win, "_ai_explain"), "Explain Current Query action should exist"
        # Clear it back out so later manual poking starts from a clean slate.
        win.ai_key_var.set("")
        win._save_ai_key()
        assert win.config.ai.get_api_key() == ""
        assert win.config.ai.enabled is False
        assert win.ai_key_status_var.get().startswith("✗"), win.ai_key_status_var.get()

        # Recent-queries history: recording is decoupled from actually
        # running a query (so a typo'd query is still recalled), and the
        # sidebar menu should offer it back for reload into the editor.
        win.config.recent_queries = []
        win.config.add_recent_query("SELECT 1")
        win.config.add_recent_query("SELECT 2 FROM dbo.Accounts")
        win._rebuild_recent_menu()
        assert win.recent_menu.index("end") == 1, "expected 2 recent-query menu entries (indices 0-1)"
        win._load_recent_query("SELECT 2 FROM dbo.Accounts")
        assert win.sql_text.text.get("1.0", "end").strip() == "SELECT 2 FROM dbo.Accounts"
        assert win.current_page == "workspace"

        # CSV export (Workspace results grid) - filedialog is modal, so
        # monkeypatch it to hand back a fixed path instead of popping a
        # real save dialog under headless Xvfb.
        import csv
        import tempfile
        from tkinter import filedialog as tkfd

        win._on_query_success(fake_result)
        root.update_idletasks()
        csv_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_export.csv")
        _orig_asksaveasfilename = tkfd.asksaveasfilename
        tkfd.asksaveasfilename = lambda **kw: csv_path
        try:
            win._export_grid_csv()
            with open(csv_path, newline="", encoding="utf-8") as f:
                exported_rows = list(csv.reader(f))
            assert exported_rows[0] == ["id", "name", "amount"]
            assert len(exported_rows) == 1 + len(fake_result.rows)

            # Excel export (Workspace results grid) - same monkeypatch trick.
            from openpyxl import load_workbook

            xlsx_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_export.xlsx")
            tkfd.asksaveasfilename = lambda **kw: xlsx_path
            win._export_grid_excel()
            wb = load_workbook(xlsx_path)
            ws = wb.active
            xlsx_rows = list(ws.iter_rows(values_only=True))
            assert xlsx_rows[0] == ("id", "name", "amount")
            assert len(xlsx_rows) == 1 + len(fake_result.rows)
            tkfd.asksaveasfilename = lambda **kw: csv_path  # restore for the next block

            # Export to Internal DB (new table) - the Workspace "Export"
            # menu's third option, reusing the same internal_store.export_
            # snapshot() path as File > Export Result to Internal DB. The
            # name prompt is a modal simpledialog, so monkeypatch it too.
            from tkinter import simpledialog as tksd
            from app.db import internal_store as internal_store_probe

            win.config.internal_db_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_export_internal.db")
            if Path(win.config.internal_db_path).exists():
                Path(win.config.internal_db_path).unlink()
            _orig_askstring = tksd.askstring
            tksd.askstring = lambda *a, **kw: "smoke_test_snapshot"
            try:
                win._export_to_internal_db()
            finally:
                tksd.askstring = _orig_askstring
            saved_snapshots = internal_store_probe.list_snapshots(win.config.internal_db_path)
            assert len(saved_snapshots) == 1, "expected exactly one new table in the internal DB"
            assert saved_snapshots[0].row_count == len(fake_result.rows)

            # Exercise the Dashboard window (separate Toplevel with embedded
            # matplotlib charts) against the same fake result, plus its own
            # inline export buttons and the snapshot Trend History tab.
            from app.ui.dashboard_window import DashboardWindow
            from app.db import internal_store

            win._open_dashboard()
            root.update_idletasks()
            root.update()
            dashboards = [w for w in root.winfo_children() if isinstance(w, DashboardWindow)]
            assert dashboards, "Dashboard window did not open"
            dash = dashboards[0]

            # KPI cards / stats rollup
            assert dash.summary == {
                "numeric_columns": 2, "categorical_columns": 1, "empty_columns": 0, "total_nulls": 0,
            }, dash.summary

            # Column search/filter on the per-column stats table
            dash.stats_filter_var.set("amount")
            dash._apply_stats_filter()
            assert len(dash.stats_tree.get_children()) == 1, "filter should narrow to just 'amount'"
            dash.stats_filter_var.set("")
            dash._apply_stats_filter()
            assert len(dash.stats_tree.get_children()) == 3, "clearing the filter should show all 3 columns"

            # Stats CSV + Excel + chart PNG export
            stats_csv_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_stats.csv")
            stats_xlsx_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_stats.xlsx")
            chart_png_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_chart.png")
            tkfd.asksaveasfilename = lambda **kw: stats_csv_path
            dash._export_stats_csv()
            assert Path(stats_csv_path).stat().st_size > 0
            tkfd.asksaveasfilename = lambda **kw: stats_xlsx_path
            dash._export_stats_excel()
            stats_wb = load_workbook(stats_xlsx_path)
            assert stats_wb.active["A1"].value == "column"
            tkfd.asksaveasfilename = lambda **kw: chart_png_path
            dash._save_figure(dash.numeric_fig)
            assert Path(chart_png_path).stat().st_size > 0

            # Correlation tab: fake_result has 2 numeric columns (id,
            # amount) so the scatter/Pearson-r panel should be built
            # rather than the "not enough numeric columns" message.
            assert hasattr(dash, "corr_fig"), "Correlation chart should be built for 2+ numeric columns"
            assert "Pearson r" in dash.corr_stat_var.get(), dash.corr_stat_var.get()
            if len(dash.numeric_cols) >= 2:
                dash.corr_x_choice.set(dash.numeric_cols[0])
                dash.corr_y_choice.set(dash.numeric_cols[1])
                dash._redraw_correlation_chart()
                assert "Pearson r" in dash.corr_stat_var.get()

            # Dashboard-wide multi-page PDF export.
            dash_pdf_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_dashboard.pdf")
            tkfd.asksaveasfilename = lambda **kw: dash_pdf_path
            dash._export_dashboard_pdf()
            assert Path(dash_pdf_path).stat().st_size > 500, "dashboard PDF looks too small/empty"

            dash.destroy()

            # Trend History tab: with no internal_db_path, there's nothing
            # to plot and it should say so rather than draw an empty chart.
            no_history_result = QueryResult(
                columns=["id"], rows=[[1]], elapsed_ms=1.0, source_table="dbo.Accounts", source_schema="dbo",
            )
            dash2 = DashboardWindow(root, no_history_result, win.config.theme, internal_db_path="")
            root.update_idletasks()
            assert dash2.trend_snapshot_count == 0
            dash2.destroy()

            # With 2+ snapshots exported for the SAME source table, the
            # trend tab should find and plot them.
            db_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_internal.db")
            if Path(db_path).exists():
                Path(db_path).unlink()
            internal_store.export_snapshot(
                db_path, "snap1", ["id"], [[1]], source_table="dbo.Accounts",
            )
            internal_store.export_snapshot(
                db_path, "snap2", ["id"], [[1], [2]], source_table="dbo.Accounts",
            )
            dash3 = DashboardWindow(root, no_history_result, win.config.theme, internal_db_path=db_path)
            root.update_idletasks()
            assert dash3.trend_snapshot_count == 2, f"expected 2 matching snapshots, got {dash3.trend_snapshot_count}"
            dash3.destroy()
        finally:
            tkfd.asksaveasfilename = _orig_asksaveasfilename

        # Results grid perf: _populate_sheet/_estimate_column_widths should
        # (a) actually size columns to their content, not leave everything
        # at the default width, and (b) stay fast even on a result well
        # past the row-sampling cap - regression check for the "running
        # queries feels laggy" fix (was: tksheet's set_all_column_widths()
        # scanning every row via slow per-cell font.measure() calls).
        from app.ui.main_window import _WIDTH_SAMPLE_CAP, _COL_MIN_WIDTH, _COL_MAX_WIDTH

        wide_col = "a_column_with_a_fairly_long_header_name"
        narrow_col = "id"
        big_rows = [
            [i, "x" * 60 if i == 0 else str(i)]
            for i in range(_WIDTH_SAMPLE_CAP + 200)  # forces the sampling branch
        ]
        t0 = time.perf_counter()
        win._populate_sheet([narrow_col, wide_col], big_rows)
        root.update_idletasks()
        elapsed_ms = (time.perf_counter() - t0) * 1000
        assert elapsed_ms < 2000, f"_populate_sheet took {elapsed_ms:.0f} ms on {len(big_rows)} rows - too slow"

        narrow_w = win.sheet.column_width(0)
        wide_w = win.sheet.column_width(1)
        assert _COL_MIN_WIDTH <= narrow_w <= _COL_MAX_WIDTH
        assert _COL_MIN_WIDTH <= wide_w <= _COL_MAX_WIDTH
        assert wide_w > narrow_w, "wide column should size larger than the narrow one"

        # ------------------------------------------------------------
        # SQL syntax highlighting: typing/inserting SQL should tag
        # keyword/string/number/comment ranges (not just leave a blob
        # of unstyled text).
        # ------------------------------------------------------------
        win.sql_text.text.delete("1.0", "end")
        win.sql_text.text.insert(
            "1.0", "SELECT id, name FROM dbo.Accounts WHERE id = 42 AND name = 'Alice' -- a comment\n"
        )
        win._highlight_sql()
        root.update_idletasks()
        assert win.sql_text.text.tag_ranges("keyword"), "SELECT/FROM/WHERE/AND should be tagged as keywords"
        assert win.sql_text.text.tag_ranges("string"), "'Alice' should be tagged as a string"
        assert win.sql_text.text.tag_ranges("number"), "42 should be tagged as a number"
        assert win.sql_text.text.tag_ranges("comment"), "-- a comment should be tagged as a comment"

        # ------------------------------------------------------------
        # Format Query button: reindents/uppercases the messy one-liner
        # left in the editor by the syntax-highlighting block above into
        # a multi-line, keyword-uppercased query - and does nothing
        # (just a status message, no crash) when the editor is blank.
        # ------------------------------------------------------------
        assert hasattr(win, "format_query_btn"), "query editor should have a Format button"
        win.sql_text.text.delete("1.0", "end")
        win.sql_text.text.insert("1.0", "select id, name from dbo.Accounts where id = 42 and name = 'Alice'")
        win._format_query()
        root.update_idletasks()
        formatted = win.sql_text.text.get("1.0", "end-1c")
        assert "SELECT" in formatted and "select" not in formatted, formatted
        assert "\n" in formatted, "prettified query should span multiple lines"
        assert "dbo.Accounts" in formatted and "42" in formatted and "Alice" in formatted, formatted

        win.sql_text.text.delete("1.0", "end")
        win._format_query()  # blank editor: should just no-op, not raise
        root.update_idletasks()
        assert win.sql_text.text.get("1.0", "end-1c") == ""

        # ------------------------------------------------------------
        # Cancel query: the generation counter should invalidate a
        # stale in-flight result instead of letting it clobber a newer
        # one, and Cancel should flip the Run/Cancel buttons back.
        # ------------------------------------------------------------
        win._query_generation = 0
        win._query_running = True
        win.run_query_btn.configure(state="disabled")
        win.cancel_query_btn.configure(state="normal")
        win.query_result = None
        stale_generation = win._query_generation
        win._query_generation += 1  # simulate a newer query having started in the meantime
        stale_result = QueryResult(columns=["x"], rows=[[1]], elapsed_ms=1.0, source_table="", source_schema="")
        win._on_query_success(stale_result, generation=stale_generation)
        assert win.query_result is None, "a stale generation's result should be discarded, not applied"
        assert win._query_running is False, "button state should reset even when the result is discarded"
        # Cancel resets state and bumps the generation so anything still
        # in flight for the old generation is now stale too.
        win._query_running = True
        gen_before_cancel = win._query_generation
        win._cancel_query()
        assert win._query_generation != gen_before_cancel, "cancel should invalidate the in-flight generation"
        assert win._query_running is False
        assert str(win.cancel_query_btn["state"]) == "disabled"

        # ------------------------------------------------------------
        # Mandatory Jira/program field: Generate Update Script should
        # refuse to run with it blank, and should stamp whatever's
        # entered into update_program once it's filled in.
        # ------------------------------------------------------------
        win._on_query_success(fake_result)
        root.update_idletasks()
        win.target_schema_var.set("dbo")
        win.target_table_var.set("Accounts")
        win.key_column_vars["id"].set(True)
        win.jira_var.set("")
        _warnings = []
        _orig_showwarning = messagebox.showwarning
        messagebox.showwarning = lambda title, msg: _warnings.append((title, msg))
        win._generate_script()
        assert _warnings, "blank Jira/program field should trigger a warning, not silently generate a script"
        messagebox.showwarning = _orig_showwarning

        win.jira_var.set("JIRA-9999")
        data = win.sheet.get_sheet_data()
        data[0][1] = "Alicia"  # was Alice - need an actual change to diff against
        win.sheet.set_sheet_data(data)
        win._generate_script()
        script_text = win.script_text.text.get("1.0", "end")
        assert "JIRA-9999" in script_text, "supplied Jira/program # should be stamped into the script"
        assert "update_program" in script_text and "update_date" in script_text and "update_user" in script_text

        # ------------------------------------------------------------
        # Tools tab: Script History - every generated script (update or
        # rollback) should show up here, auto-refreshed just by opening
        # the Tools page, and be viewable in full via the popup.
        # ------------------------------------------------------------
        from tkinter import simpledialog as tksd2

        win._show_page("tools")
        root.update_idletasks()
        assert hasattr(win, "script_history_tree"), "Tools tab should have a Script History section"
        history_rows = win.script_history_tree.get_children()
        # win.config.internal_db_path was redirected to a fresh temp file
        # earlier (Export to Internal DB block above) and the blank-Jira
        # attempt just above never got as far as generate_update_script -
        # so exactly the one JIRA-9999 generation should be here.
        assert len(history_rows) == 1, f"expected exactly 1 history entry, got {len(history_rows)}"
        latest = win._script_history_cache[0]
        assert latest.kind == "update"
        assert latest.table_name == "Accounts"
        assert latest.program == "JIRA-9999"
        assert "JIRA-9999" in latest.sql_text
        assert latest.source == "desktop"

        win.script_history_tree.selection_set(str(latest.id))
        win._view_selected_history_entry()
        popup_windows = [w for w in root.winfo_children() if isinstance(w, tb.Toplevel) and w.winfo_exists()]
        assert popup_windows, "View Full Script should open a popup window"
        for w in popup_windows:
            w.destroy()

        win.config.saved_queries = []
        win.sql_text.text.delete("1.0", "end")
        win.sql_text.text.insert("1.0", "SELECT * FROM dbo.Widgets")
        _orig_askstring2 = tksd2.askstring
        tksd2.askstring = lambda *a, **kw: "My Widgets Query"
        try:
            win._save_current_query()
        finally:
            tksd2.askstring = _orig_askstring2
        assert [q.name for q in win.config.saved_queries] == ["My Widgets Query"]
        assert len(win.saved_queries_tree.get_children()) == 1

        win.saved_queries_tree.selection_set(win.saved_queries_tree.get_children()[0])
        win.sql_text.text.delete("1.0", "end")  # clear so we can tell Load actually did something
        win._load_selected_saved_query()
        assert win.sql_text.text.get("1.0", "end").strip() == "SELECT * FROM dbo.Widgets"
        assert win.current_page == "workspace"

        win._show_page("tools")
        win.saved_queries_tree.selection_set(win.saved_queries_tree.get_children()[0])
        win._delete_selected_saved_query()
        assert win.config.saved_queries == []
        assert len(win.saved_queries_tree.get_children()) == 0

        # ------------------------------------------------------------
        # Tools tab: Schema Validation - exercises the render path
        # (_on_schema_validated) directly, since the actual lookup needs
        # a live DB connection; app.core.schema_check's own logic is
        # covered by tests/test_schema_check.py.
        # ------------------------------------------------------------
        win._on_schema_validated(
            schema_check.validate_columns(["id", "amount"], ["id"], {"id": "int", "amount": "money"}),
            "dbo", "Accounts",
        )
        assert win.schema_validation_var.get().startswith("✓"), win.schema_validation_var.get()
        win._on_schema_validated(
            schema_check.validate_columns(["id", "ghost_col"], ["id"], {"id": "int"}),
            "dbo", "Accounts",
        )
        assert win.schema_validation_var.get().startswith("✗") and "ghost_col" in win.schema_validation_var.get()

        # ------------------------------------------------------------
        # Tools tab: Rollback Script - reverses the forward script
        # generated above (JIRA-9999, name changed Alice -> Alicia).
        # ------------------------------------------------------------
        win._generate_rollback_script()
        rollback_text = win.rollback_output.text.get("1.0", "end")
        assert "Rollback" in rollback_text
        assert "Alice" in rollback_text, "rollback should restore the original 'Alice' value"
        assert "JIRA-9999" in rollback_text, "rollback should re-stamp the same program #"

        # ------------------------------------------------------------
        # Tools tab: Snapshot Diff Viewer - export two snapshots that
        # differ by one row, then compare them.
        # ------------------------------------------------------------
        win.config.internal_db_path = str(Path(tempfile.gettempdir()) / "scriptgen_smoke_diff.db")
        if Path(win.config.internal_db_path).exists():
            Path(win.config.internal_db_path).unlink()
        internal_store.export_snapshot(
            win.config.internal_db_path, "diff_older", ["id", "amount"], [[1, 10], [2, 20]], source_table="dbo.Widgets",
        )
        internal_store.export_snapshot(
            win.config.internal_db_path, "diff_newer", ["id", "amount"], [[1, 10], [2, 25], [3, 30]], source_table="dbo.Widgets",
        )
        win._refresh_diff_snapshot_list()
        assert len(win._diff_snapshots_cache) == 2
        # list_snapshots() orders newest-first, so index 0 is diff_newer
        # and index 1 is diff_older - Older/Newer combos need it that
        # way round for "1 added" (row 3) to come out as added, not
        # removed.
        older_idx = next(i for i, s in enumerate(win._diff_snapshots_cache) if s.display_name == "diff_older")
        newer_idx = next(i for i, s in enumerate(win._diff_snapshots_cache) if s.display_name == "diff_newer")
        win.diff_a_combo.current(older_idx)
        win.diff_b_combo.current(newer_idx)
        win.diff_key_var.set("id")
        win._compare_snapshots()
        assert "1 added" in win.diff_summary_var.get(), win.diff_summary_var.get()
        assert "1 changed" in win.diff_summary_var.get(), win.diff_summary_var.get()
        assert len(win.diff_tree.get_children()) == 2  # 1 added + 1 changed row rendered (unchanged rows aren't)

        # ------------------------------------------------------------
        # Tools tab: Environment switcher - add a second connection and
        # switch to it; Settings dialog changes should keep this combo
        # in sync too (exercised via _refresh_env_combo directly, since
        # opening the modal ConfigDialog itself is covered separately).
        # ------------------------------------------------------------
        from app.config import ConnectionConfig as _ConnectionConfig

        win.config.add_connection("staging", _ConnectionConfig(name="Staging"))
        win._refresh_env_combo()
        assert "staging" in win._env_keys
        staging_idx = win._env_keys.index("staging")
        win.env_combo.current(staging_idx)
        win._switch_environment()
        assert win.config.active_connection == "staging"
        # switch back so later manual poking starts from the default connection
        default_idx = win._env_keys.index("default") if "default" in win._env_keys else 0
        win.env_combo.current(default_idx)
        win._switch_environment()

        # ------------------------------------------------------------
        # AI Assist: NL -> WHERE clause builder and the Script page's
        # AI Review button. First the offline "no key configured"
        # short-circuit (must not crash, must not spawn a network call);
        # then a full round trip with ai_assist's Gemini calls
        # monkeypatched so this stays a fast, offline smoke test.
        # ------------------------------------------------------------
        win.config.ai.enabled = False
        win.config.ai.api_key_enc = ""
        _infos = []
        _orig_showinfo = messagebox.showinfo
        messagebox.showinfo = lambda title, msg: _infos.append((title, msg))
        win._ai_build_where()
        win._ai_review_script()
        messagebox.showinfo = _orig_showinfo
        assert len(_infos) == 2, "both AI actions should short-circuit with a dialog when no key is configured"

        # These AI actions run their Gemini call on a background thread
        # and hand the result back via root.after(0, ...). That's the
        # right pattern for the real app (which runs a genuine
        # root.mainloop()), but Tk refuses cross-thread root.after()
        # calls when nothing is actually inside mainloop() - which is
        # exactly this headless test's situation (it drives Tk via
        # repeated root.update() instead). Rather than fight Tk's
        # threading model, run the worker synchronously on the main
        # thread here so root.after() is called from a thread Tk
        # accepts - still exercises the exact same worker/callback code.
        import threading as _threading_mod

        class _SyncThread:
            def __init__(self, target=None, daemon=None, **kw):
                self._target = target

            def start(self):
                self._target()

        win.config.ai.set_api_key("fake-smoke-key")
        _orig_suggest_where = ai_assist.suggest_where_clause
        _orig_review_script = ai_assist.review_script
        _orig_thread_cls = _threading_mod.Thread
        ai_assist.suggest_where_clause = lambda *a, **kw: ai_assist.AIResponse(
            success=True, text="```sql\nstatus = 'active'\n```", extracted_sql="status = 'active'"
        )
        ai_assist.review_script = lambda *a, **kw: ai_assist.AIResponse(success=True, text="Looks safe.")
        _threading_mod.Thread = _SyncThread
        try:
            win.ai_where_request_var.set("only active rows")
            win._ai_build_where()
            root.update()
            assert win._last_ai_where == "status = 'active'", win.ai_where_output.text.get("1.0", "end")
            win._append_ai_where_to_query()
            assert "WHERE status = 'active'" in win.sql_text.text.get("1.0", "end")

            win._ai_review_script()
            root.update()
            assert "Looks safe" in win.ai_review_output.text.get("1.0", "end")
        finally:
            ai_assist.suggest_where_clause = _orig_suggest_where
            ai_assist.review_script = _orig_review_script
            _threading_mod.Thread = _orig_thread_cls
            win.config.ai.set_api_key("")

        # Restore a normal result so later manual poking around isn't confusing.
        win._on_query_success(fake_result)
        root.update_idletasks()
        root.update()

        # Exercise the light/dark theme toggle
        assert win.config.theme == LIGHT_THEME
        win._toggle_theme()
        assert win.config.theme == DARK_THEME
        root.update_idletasks()
        root.update()

        print(
            "SMOKE TEST PASSED: MainWindow + Dashboard construct and theme correctly; grid/script/rollback/"
            "diff/schema-check/AI/env-switcher/syntax-highlight/cancel-query flows all work."
        )
    except Exception:
        print("SMOKE TEST FAILED:")
        traceback.print_exc()
        sys.exit(1)
    finally:
        root.destroy()


if __name__ == "__main__":
    main()
